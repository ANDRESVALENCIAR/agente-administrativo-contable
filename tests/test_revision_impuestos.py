"""Pruebas de la revisión de impuestos con datos 100 % sintéticos (nunca archivos reales)."""
import os
import sys
import tempfile
from datetime import date, datetime

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ["DATABASE_PATH"] = os.path.join(tempfile.mkdtemp(), "rev.db")

import openpyxl  # noqa: E402
import pytest  # noqa: E402
from reportlab.pdfgen import canvas  # noqa: E402

from modulos import revision_350 as r350  # noqa: E402
from modulos.revision_lectores import leer_excel_formulario, leer_pdf_formulario  # noqa: E402

CAMPOS = {77: 10_000_000, 93: 300_000, 36: 2_000_000, 49: 50_000, 59: 100_000_000, 68: 1_200_000,
          130: 1_550_000, 131: 30_000, 134: 30_000, 136: 1_580_000, 138: 1_580_000}


def _pdf(ruta, campos, con_980=False):
    c = canvas.Canvas(str(ruta), pagesize=(612, 792))
    y = 700
    for k, v in campos.items():
        c.drawString(150, y, str(k))
        c.drawString(250, y + 2, f"{v:,.0f}")
        y -= 20
    if con_980:
        c.drawString(150, y, "980")
        c.drawString(250, y + 2, f"{campos[138]:,.0f}")
    c.save()


def _libro(ruta, campos):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Datos"
    for f in [("DATOS", "DESCRIPCIÓN"), ("Periodo", 2026), ("Mes", 9), ("Actividad", 4659)]:
        ws.append(f)
    pr = wb.create_sheet("Prueba")
    pr.append(["EMPRESA"])
    pr.append([None, "GRUPO", "DESCRIPCION", "SALDO ANTERIOR", "DEBITOS", "CREDITOS", "NUEVO SALDO"])
    for fila in [(2, "41", "INGRESOS", 0, 0, 100_000_000, 0), (4, "2365", "RETEFUENTE", 0, 0, 1_550_000, 0),
                 (6, "236505", "SALARIOS", 0, 0, 300_000, 0), (6, "236540", "COMPRAS", 0, 0, 50_000, 0),
                 (6, "236575", "AUTORRET", 0, 0, 1_200_000, 0), (4, "2367", "RETEIVA", 0, 0, 30_000, 0),
                 (8, "23658001", "RETEFUENTE X PAGAR", 2_000_000, 0, 0, 2_000_000)]:
        pr.append(list(fila))
    ax = wb.create_sheet("Auxiliar")
    ax.append(["SUB", "CUENTA", "DESCRIPCION", "NIT", "NOMBRE", "COMPROBANTE", "FECHA", "DETALLE", "SALDO"])
    ax.append(["2365", "23658001", "", 1, "EIF", "NC 1", datetime(2026, 6, 30), "CIERRE DE IMPUESTOS JUNIO", -1_000_000])
    ax.append(["2365", "23658001", "", 2, "DIAN", "CE 2", datetime(2026, 7, 10), "PAGO RETENCION JUNIO", 1_000_000])
    ax.append(["2365", "23658001", "", 1, "EIF", "NC 3", datetime(2026, 7, 31), "CIERRE DE IMPUESTO JULIO", -2_000_000])
    ax.append(["2367", "23670501", "", 9, "PROVEEDOR SIMPLE Nit: 9", "FC 4", datetime(2026, 9, 5), "COMPRA", -30_000])
    rt = wb.create_sheet("RetenciónHT")
    rt.append(["MES", "TP", "CUENTA", "DESCRIPCION", "NIT", "NOMBRE", "DETALLE", "INV-CRUC-BASE", "SALDO"])
    rt.append(["Sep", "PJ", "23652501", "SERVICIOS NO DECLARANTES 6%", "800", "EMPRESA SAS Nit: 800", "x", 1_000_000, -60_000])
    rt.append(["Sep", "PJ", "23654001", "COMPRAS DECLARANTES 2.5%", "801", "OTRA SAS", "x", 2_000_000, -50_000])
    au = wb.create_sheet("AutorrentaHT")
    au.append(["MES", "TP", "SALDO"])
    au.append(["Sep", "PJ", -1_200_000])
    fo = wb.create_sheet("Formato 2026 Persona Uno")
    for f in [("UVT 2026", 52374), ("Totalidad de ingresos renta de trabajo", 12_000_000),
              ("Ingreso Laboral Mensual Base para Retención en la Fuente", 10_000_000),
              ("RETENCION EN LA FUENTE POR ART 383", 0)]:
        fo.append(f)
    # base 10.000.000 = 190,93 UVT → (190,93−150)×28% + 10 = 21,46 UVT = $1.124.000
    fo.cell(row=4, column=2, value=1_124_000)
    fm = wb.create_sheet("Form 350")
    for k, v in campos.items():
        fm.append([None, k, v])
    wb.save(ruta)


@pytest.fixture
def insumos(tmp_path):
    pdf, xl = tmp_path / "borrador.pdf", tmp_path / "liq.xlsx"
    _pdf(pdf, CAMPOS)
    _libro(xl, CAMPOS)
    return pdf, xl


def test_lee_pdf_y_excel(insumos):
    pdf, xl = insumos
    p = leer_pdf_formulario(pdf)["campos"]
    assert p[77] == 10_000_000 and p[138] == 1_580_000
    x = leer_excel_formulario(xl, "350")
    assert x["hoja"] == "Form 350" and x["campos"][130] == 1_550_000


def test_tabla_383():
    assert r350._tabla_383(90) == 0
    assert round(r350._tabla_383(146.29), 4) == round((146.29 - 95) * 0.19, 4)
    assert round(r350._tabla_383(200), 4) == round((200 - 150) * 0.28 + 10, 4)


def test_aritmetica_detecta_error():
    malo = {**CAMPOS, 138: 1_000_000}
    _, h = r350.aritmetica(malo)
    assert any("138" in x["titulo"] for x in h)
    ok, h2 = r350.aritmetica(CAMPOS)
    assert not [x for x in h2 if x["nivel"] == "CRITICO"]


def test_revision_completa(insumos):
    pdf, xl = insumos
    res = r350.revisar(leer_pdf_formulario(pdf)["campos"], str(xl), 2026, 9, hoy=date(2026, 10, 10),
                       perfil_reteica=None)
    titulos = " | ".join(h["titulo"] for h in res["hallazgos"])
    assert "980" in titulos                              # pago total vacío
    assert "persona jurídica con tarifa de NO declarante" in titulos
    assert "Julio 2026 SIN PAGAR" in titulos           # cierre de julio sin pago
    assert "régimen SIMPLE" in titulos
    assert not any(h["area"] == "Formulario" and h["nivel"] == "CRITICO" for h in res["hallazgos"])
    assert res["tablas"]["art_383"][0]["ok"] is True    # 190,93 UVT → $1.124.000


def test_estado_pago():
    v = date(2026, 8, 18)
    assert r350._estado_pago(v, date(2026, 8, 10), date(2026, 10, 1))[0] == "OK"
    assert r350._estado_pago(v, date(2026, 9, 1), date(2026, 10, 1))[0] == "MEDIO"
    assert r350._estado_pago(v, date(2026, 11, 1), date(2026, 12, 1))[0] == "CRITICO"
    assert r350._estado_pago(v, None, date(2026, 8, 1))[0] == "INFO"


def test_guardado_y_historial(insumos, tmp_path, monkeypatch):
    from config import cfg
    from database import inicializar_db
    from modulos import revision_impuestos as ri

    monkeypatch.setattr(cfg, "DATABASE_PATH", str(tmp_path / "rev.db"))

    monkeypatch.setattr(ri, "RAIZ", tmp_path / "rev")
    inicializar_db()
    pdf, xl = insumos
    rev = ri.revisar_rutas("RET", "2026-09", [str(pdf), str(xl)], hoy=date(2026, 10, 10))
    assert rev["veredicto"] and rev["informe_path"].endswith(".html")
    assert "Retención en la fuente" in open(rev["informe_path"], encoding="utf-8").read()
    assert ("RET", "2026-09") in ri.ultima_por_periodo(2026)
    assert ri.periodo_sugerido("RET", date(2026, 10, 10)) == "2026-09"
    assert ri.periodo_sugerido("IVA", date(2026, 10, 10)) == "2026-B4"
