"""Pruebas del control de pagos (libro Maestro + Calendario Anual) y alertas 24h."""
import os
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

_TMP = tempfile.mkdtemp()
os.environ["DATABASE_PATH"] = os.path.join(_TMP, "cp.db")
os.environ["PAGOS_CONTROL_EXCEL"] = os.path.join(_TMP, "control.xlsx")
os.environ["PAGOS_ALERTA_EMAILS"] = "pagos@empresa.com"
os.environ["PAGOS_ALERTA_WHATSAPP"] = "+573001234567"
os.environ["APP_PUBLIC_URL"] = "https://shaki.example.com"
os.environ["PAGOS_LINK_SECRET"] = "secreto-de-prueba"

from config import cfg  # noqa: E402

cfg.DATABASE_PATH = os.environ["DATABASE_PATH"]
cfg.PAGOS_CONTROL_EXCEL = os.environ["PAGOS_CONTROL_EXCEL"]
cfg.PAGOS_ALERTA_EMAILS = os.environ["PAGOS_ALERTA_EMAILS"]
cfg.PAGOS_ALERTA_WHATSAPP = os.environ["PAGOS_ALERTA_WHATSAPP"]
cfg.APP_PUBLIC_URL = os.environ["APP_PUBLIC_URL"]
cfg.PAGOS_LINK_SECRET = os.environ["PAGOS_LINK_SECRET"]
cfg.PAGOS_ESCRIBIR_EXCEL = True

from database import inicializar_db  # noqa: E402
from modulos import control_pagos as cp  # noqa: E402


def _libro(path: str) -> None:
    """Libro mínimo con la misma estructura que 'Control de pagos para Agente Shaki'."""
    from openpyxl import Workbook

    wb = Workbook()
    ins = wb.active
    ins.title = "Instrucciones"
    ins["A3"], ins["B3"] = "Año del calendario", 2026
    m = wb.create_sheet("Maestro de Pagos")
    m.append(["ID", "Detalle", "Tipo", "Proveedor", "Referencia / Cuenta", "Centro de costos",
              "Periodicidad", "Día del mes", "Día de la semana", "Prioridad"])
    m.append(["P01", "ARRIENDO", "PAGO", "INMOBILIARIA", None, "EIF", "MENSUAL", 18, None, 1])
    m.append(["P02", "NÓMINA", "PAGO", "PERSONAL", None, "EIF", "MENSUAL", 30, None, 1])
    m.append(["P03", "GRABAR PLANILLA", "TAREA", "MI PLANILLA", None, "SANDRA", "MENSUAL", 13, None, 2])
    m.append(["S01", "NÓMINA SEMANAL", "PAGO", "EIF", None, "EIF", "SEMANAL", None, 5, 1])
    c = wb.create_sheet("Calendario Anual")
    c.append(["ID", "Ocurrencia", "Periodicidad", "Fecha límite", "Mes", "Nombre mes", "Día semana",
              "Fecha sugerida de pago", "Prioridad", "Tipo", "Detalle", "Proveedor", "Referencia / Cuenta",
              "Centro de costos", "Valor", "Estado", "Fecha pago real", "Observaciones", "Alerta"])
    for mes in range(1, 13):
        c.append(["P01", mes] + [None] * 12 + [None, "PAGADO" if mes < 10 else "PENDIENTE"])
        c.append(["P02", mes] + [None] * 13 + ["PENDIENTE"])
        c.append(["P03", mes] + [None] * 13 + ["PENDIENTE"])
    for sem in range(1, 53):
        c.append(["S01", sem] + [None] * 13 + ["PENDIENTE"])
    f = wb.create_sheet("Festivos")
    f.append(["Fecha", "Festivo"])
    f.append([date(2026, 10, 12), "Día de la Raza"])
    wb.save(path)


class TestControlPagos(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        inicializar_db()
        _libro(cfg.PAGOS_CONTROL_EXCEL)
        cls.res = cp.importar_excel()

    def test_importa_todas_las_cuotas(self):
        self.assertEqual(self.res["total"], 12 * 3 + 52)
        self.assertEqual(cp.importar_excel()["nuevos"], 0)  # idempotente

    def test_fechas_reglas_del_libro(self):
        # Febrero no tiene 30: último día del mes
        self.assertEqual(cp.fecha_limite("MENSUAL", 2, 2026, 30, None), date(2026, 2, 28))
        # Semanal viernes: primera semana de 2026 = viernes 2 de enero
        self.assertEqual(cp.fecha_limite("SEMANAL", 1, 2026, None, 5), date(2026, 1, 2))
        # Domingo 18 oct 2026 se adelanta al viernes 16
        self.assertEqual(cp.fecha_sugerida(date(2026, 10, 18), set()), date(2026, 10, 16))
        # Martes 13 oct es hábil; lunes 12 es festivo → el 12 se adelanta al viernes 9
        self.assertEqual(cp.fecha_sugerida(date(2026, 10, 12), {date(2026, 10, 12)}), date(2026, 10, 9))

    def test_alerta_24h_una_sola_vez(self):
        pendientes = cp.pagos_para_alerta(date(2026, 10, 12))  # mañana = martes 13
        self.assertEqual([p["id_pago"] for p in pendientes], ["P03"])
        r = cp.enviar_alertas_24h(date(2026, 10, 12))
        self.assertEqual(r["pagos"], 1)
        self.assertEqual(r["correos"], 1)
        self.assertEqual(r["whatsapp"], 1)
        self.assertEqual(cp.enviar_alertas_24h(date(2026, 10, 12))["pagos"], 0)

    def test_mensajes_dicen_pendiente_y_traen_boton(self):
        p = cp.listar(date(2026, 10, 16), date(2026, 10, 16), cp.ABIERTOS)[0]
        html = cp.html_correo([p])
        self.assertIn("PENDIENTE DE PAGO", html)
        self.assertIn("https://shaki.example.com/?pagar=", html)
        self.assertIn("PENDIENTE DE PAGO", cp.texto_whatsapp(p))
        self.assertIn("?pagar=", cp.texto_whatsapp(p))

    def test_token_firmado(self):
        tok = cp.generar_token(5)
        self.assertEqual(cp.verificar_token(tok), 5)
        self.assertIsNone(cp.verificar_token("5.falso"))
        self.assertIsNone(cp.verificar_token(tok.replace("5.", "6.", 1)))
        self.assertIsNone(cp.verificar_token("basura"))

    def test_ya_lo_pague_actualiza_bd_y_excel(self):
        from openpyxl import load_workbook

        p = cp.listar(date(2026, 10, 16), date(2026, 10, 16), cp.ABIERTOS)[0]
        pid = cp.verificar_token(cp.generar_token(p["id"]))
        r = cp.marcar_pagado(pid, por="enlace de alerta", referencia="TRX-1")
        self.assertTrue(r["ok"])
        self.assertTrue(r["excel"])
        self.assertTrue(cp.marcar_pagado(pid)["ya_estaba"])
        self.assertEqual(cp.obtener(pid)["estado"], "PAGADO")
        ws = load_workbook(cfg.PAGOS_CONTROL_EXCEL)["Calendario Anual"]
        self.assertEqual(ws.cell(p["excel_fila"], 16).value, "PAGADO")
        self.assertIn("TRX-1", ws.cell(p["excel_fila"], 18).value)
        self.assertTrue(Path(cfg.PAGOS_CONTROL_EXCEL + ".bak").exists())
        # Releer el Excel no deshace el pago
        cp.importar_excel()
        self.assertEqual(cp.obtener(pid)["estado"], "PAGADO")

    def test_semaforo(self):
        base = {"estado": "PENDIENTE"}
        hoy = date(2026, 10, 10)
        self.assertEqual(cp.semaforo({**base, "fecha_sugerida": "2026-10-09"}, hoy), "VENCIDO")
        self.assertEqual(cp.semaforo({**base, "fecha_sugerida": "2026-10-10"}, hoy), "HOY")
        self.assertEqual(cp.semaforo({**base, "fecha_sugerida": "2026-10-11"}, hoy), "MAÑANA")
        self.assertEqual(cp.semaforo({**base, "fecha_sugerida": "2026-10-13"}, hoy), "PRÓXIMO")
        self.assertEqual(cp.semaforo({**base, "fecha_sugerida": "2026-10-30"}, hoy), "AL DÍA")
        self.assertEqual(cp.semaforo({"estado": "PAGADO", "fecha_sugerida": "2026-01-01"}, hoy), "PAGADO")


if __name__ == "__main__":
    unittest.main()
