"""Sincronización bidireccional CONTROL_DE_PAGOS.xlsx ↔ SQLite."""
from __future__ import annotations

import logging
import os
import sqlite3
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from config import cfg
from database import get_conn, registrar_accion

logger = logging.getLogger(__name__)

SHEET_NAME = "CONTROL DE PAGOS"
COLUMNAS = [
    "prioridad",
    "dia_semana",
    "fecha_limite",
    "detalle",
    "proveedor",
    "periodo",
    "anio",
    "centro_costos",
    "valor_a_pagar",
    "valor_pagado",
    "fecha_pago",
    "cuenta_origen_destino",
    "pago_total",
    "observaciones",
    "periodicidad",
]

CENTROS_COSTO = [
    "EIF",
    "SANDRA",
    "ANDRES",
    "SYA",
    "EIF-HUGO",
    "ELECTROMATICA",
    "CBDO",
    "CONFINAT",
    "IAI",
    "EIF Y ELECTROMATICA",
    "NUBANK",
    "BRITTO",
    "CUEROS VELEZ",
    "MARTIN AGUILERA",
    "ANDRES JUAN",
    "JUANES",
    "JUAN ESTEBAN",
    "NIÑOS",
    "USB",
]

PERIODICIDADES = [
    "MENSUAL",
    "ANUAL",
    "IMPUESTO",
    "IMPUESTO BIMESTRAL",
    "IMPUESTO MENSUAL",
    "IMPUESTO ANUAL",
    "SEMESTRAL",
    "ESPORADICO",
    "COMPRA",
    "REUNION",
    "CITA MEDICA",
    "USB",
]


def ruta_excel_pagos() -> Path:
    """Resuelve la ruta del CONTROL DE PAGOS (env, config, raíz o data/)."""
    candidatos: list[Path] = []
    env = os.getenv("CONTROL_DE_PAGOS", "").strip()
    if env:
        candidatos.append(Path(env))
    cfg_path = getattr(cfg, "CONTROL_DE_PAGOS", None)
    if cfg_path:
        candidatos.append(Path(cfg_path))
    base = Path(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    candidatos.extend(
        [
            base / "CONTROL_DE_PAGOS.xlsx",
            base / "CONTROL DE PAGOS.xlsx",
            base / "data" / "CONTROL_DE_PAGOS.xlsx",
            base / "data" / "CONTROL DE PAGOS.xlsx",
            Path(os.path.dirname(base))
            / "Agente ADmin Shaki"
            / "CONTROL DE PAGOS (1).xlsx",
            Path(os.path.dirname(base)) / "Agente ADmin Shaki" / "CONTROL DE PAGOS.xlsx",
        ]
    )
    for p in candidatos:
        if p.is_file():
            return p
    return candidatos[0] if candidatos else base / "CONTROL_DE_PAGOS.xlsx"


def calcular_estado(fecha_limite: date | None, fecha_pago: Any, estado_actual: str | None = None) -> str:
    """Calcula estado operativo del pago según fechas (respeta APROBADO/RECHAZADO)."""
    if fecha_pago not in (None, "", pd.NaT) and not (isinstance(fecha_pago, float) and pd.isna(fecha_pago)):
        return "PAGADO"
    if estado_actual in ("APROBADO", "RECHAZADO"):
        return estado_actual
    if not fecha_limite:
        return "PENDIENTE"
    hoy = date.today()
    if isinstance(fecha_limite, datetime):
        fecha_limite = fecha_limite.date()
    diff = (fecha_limite - hoy).days
    if diff < 0:
        return "VENCIDO"
    if diff == 0:
        return "HOY"
    if diff <= 3:
        return "URGENTE"
    return "PENDIENTE"


def inicializar_tablas_pagos() -> None:
    """Crea tablas del módulo PAGOS (control Excel + aprobaciones + log + checklist)."""
    conn = get_conn()
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS pagos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            excel_fila INTEGER,
            prioridad INTEGER DEFAULT 1,
            dia_semana TEXT,
            fecha_limite DATE NOT NULL,
            detalle TEXT NOT NULL,
            proveedor TEXT,
            periodo TEXT,
            anio INTEGER,
            centro_costos TEXT,
            valor_a_pagar REAL,
            valor_pagado REAL,
            fecha_pago DATE,
            cuenta_origen_destino TEXT,
            pago_total TEXT,
            observaciones TEXT,
            periodicidad TEXT,
            estado TEXT DEFAULT 'PENDIENTE',
            creado_en TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            actualizado_en TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            origen TEXT DEFAULT 'EXCEL',
            sync_hash TEXT,
            UNIQUE(excel_fila)
        );

        CREATE TABLE IF NOT EXISTS pagos_aprobaciones (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            pago_id INTEGER REFERENCES pagos(id),
            accion TEXT NOT NULL,
            usuario TEXT DEFAULT 'ADMIN',
            fecha_accion TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            comentario TEXT
        );

        CREATE TABLE IF NOT EXISTS pagos_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            pago_id INTEGER REFERENCES pagos(id),
            campo_modificado TEXT,
            valor_anterior TEXT,
            valor_nuevo TEXT,
            modificado_por TEXT DEFAULT 'SISTEMA',
            fecha_modificacion TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS pagos_sync_meta (
            id INTEGER PRIMARY KEY CHECK (id = 1),
            ultima_importacion TIMESTAMP,
            ultima_exportacion TIMESTAMP,
            registros INTEGER DEFAULT 0
        );

        CREATE INDEX IF NOT EXISTS idx_pagos_estado ON pagos(estado);
        CREATE INDEX IF NOT EXISTS idx_pagos_fecha ON pagos(fecha_limite);
        CREATE INDEX IF NOT EXISTS idx_pagos_proveedor ON pagos(proveedor);
        CREATE INDEX IF NOT EXISTS idx_pagos_centro ON pagos(centro_costos);
        """
    )
    # Ampliar checklist existente
    cols = {r[1] for r in conn.execute("PRAGMA table_info(checklist_pagos_periodo)").fetchall()}
    for col, ddl in [
        ("aportes_ss_pagados", "INTEGER DEFAULT 0"),
        ("retencion_fuente_calculada", "INTEGER DEFAULT 0"),
    ]:
        if col not in cols:
            try:
                conn.execute(f"ALTER TABLE checklist_pagos_periodo ADD COLUMN {col} {ddl}")
            except sqlite3.OperationalError:
                pass
    conn.commit()
    conn.close()


def _to_date(val: Any) -> date | None:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    if isinstance(val, datetime):
        return val.date()
    if isinstance(val, date):
        return val
    if isinstance(val, str):
        s = val.strip()
        if not s:
            return None
        try:
            return date.fromisoformat(s[:10])
        except ValueError:
            pass
    try:
        ts = pd.to_datetime(val, errors="coerce", dayfirst=True)
        if pd.isna(ts):
            return None
        return ts.date()
    except Exception:
        return None


def _to_float(val: Any) -> float | None:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    if isinstance(val, str):
        s = val.strip().replace("$", "").replace(" ", "").replace(".", "").replace(",", ".")
        if not s:
            return None
        try:
            return float(s)
        except ValueError:
            return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def _to_int(val: Any) -> int | None:
    f = _to_float(val)
    if f is None:
        return None
    return int(f)


def _txt(val: Any) -> str | None:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return None
    s = str(val).strip()
    return s or None


def _hash_fila(row: dict[str, Any]) -> str:
    parts = [
        str(row.get("fecha_limite") or ""),
        str(row.get("detalle") or ""),
        str(row.get("proveedor") or ""),
        str(row.get("valor_a_pagar") or ""),
        str(row.get("valor_pagado") or ""),
        str(row.get("fecha_pago") or ""),
        str(row.get("centro_costos") or ""),
        str(row.get("periodicidad") or ""),
        str(row.get("observaciones") or ""),
    ]
    return "|".join(parts)


def leer_excel_pagos(path: Path | None = None) -> pd.DataFrame:
    """Lee la hoja CONTROL DE PAGOS con encabezado en fila 2."""
    path = path or ruta_excel_pagos()
    if not path.is_file():
        raise FileNotFoundError(f"No se encontró CONTROL DE PAGOS: {path}")
    df = pd.read_excel(
        path,
        sheet_name=SHEET_NAME,
        header=1,
        usecols="A:O",
        names=COLUMNAS,
    )
    return df


def importar_excel_a_sqlite(path: Path | None = None) -> dict[str, int]:
    """
    Importa CONTROL_DE_PAGOS.xlsx a SQLite.
    Detecta filas nuevas/actualizadas por excel_fila (fila física en Excel).
    No sobrescribe pagos MANUAL ni estados APROBADO/RECHAZADO/PAGADO del dashboard
    salvo que el Excel traiga fecha_pago.
    """
    inicializar_tablas_pagos()
    path = path or ruta_excel_pagos()
    df = leer_excel_pagos(path)

    nuevos = 0
    actualizados = 0
    procesados = 0
    conn = get_conn()
    c = conn.cursor()

    for idx, raw in df.iterrows():
        excel_fila = int(idx) + 3  # header=1 → primera data = fila Excel 3
        detalle = _txt(raw.get("detalle"))
        fecha_limite = _to_date(raw.get("fecha_limite"))
        if not detalle or not fecha_limite:
            continue

        fecha_pago = _to_date(raw.get("fecha_pago"))
        row = {
            "prioridad": _to_int(raw.get("prioridad")) or 1,
            "dia_semana": _txt(raw.get("dia_semana")),
            "fecha_limite": fecha_limite.isoformat(),
            "detalle": detalle,
            "proveedor": _txt(raw.get("proveedor")),
            "periodo": _txt(raw.get("periodo")),
            "anio": _to_int(raw.get("anio")),
            "centro_costos": _txt(raw.get("centro_costos")),
            "valor_a_pagar": _to_float(raw.get("valor_a_pagar")),
            "valor_pagado": _to_float(raw.get("valor_pagado")),
            "fecha_pago": fecha_pago.isoformat() if fecha_pago else None,
            "cuenta_origen_destino": _txt(raw.get("cuenta_origen_destino")),
            "pago_total": _txt(raw.get("pago_total")),
            "observaciones": _txt(raw.get("observaciones")),
            "periodicidad": _txt(raw.get("periodicidad")),
        }
        sync_hash = _hash_fila(row)
        procesados += 1

        c.execute(
            "SELECT id, sync_hash, estado, origen, fecha_pago FROM pagos WHERE excel_fila=?",
            (excel_fila,),
        )
        existing = c.fetchone()
        estado_prev = existing[2] if existing else None
        estado = calcular_estado(fecha_limite, fecha_pago, estado_prev)

        if existing is None:
            c.execute(
                """INSERT INTO pagos (
                    excel_fila, prioridad, dia_semana, fecha_limite, detalle, proveedor,
                    periodo, anio, centro_costos, valor_a_pagar, valor_pagado, fecha_pago,
                    cuenta_origen_destino, pago_total, observaciones, periodicidad,
                    estado, origen, sync_hash, actualizado_en
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,CURRENT_TIMESTAMP)""",
                (
                    excel_fila,
                    row["prioridad"],
                    row["dia_semana"],
                    row["fecha_limite"],
                    row["detalle"],
                    row["proveedor"],
                    row["periodo"],
                    row["anio"],
                    row["centro_costos"],
                    row["valor_a_pagar"],
                    row["valor_pagado"],
                    row["fecha_pago"],
                    row["cuenta_origen_destino"],
                    row["pago_total"],
                    row["observaciones"],
                    row["periodicidad"],
                    estado,
                    "EXCEL",
                    sync_hash,
                ),
            )
            nuevos += 1
        else:
            pago_id, old_hash, old_estado, origen, old_fp = existing
            if origen == "MANUAL" and not fecha_pago:
                # Solo refrescar estado por fechas si no está pagado
                nuevo_est = calcular_estado(fecha_limite, old_fp, old_estado)
                if nuevo_est != old_estado:
                    c.execute(
                        "UPDATE pagos SET estado=?, actualizado_en=CURRENT_TIMESTAMP WHERE id=?",
                        (nuevo_est, pago_id),
                    )
                    actualizados += 1
                continue
            if old_hash == sync_hash:
                # Recalcular estado por calendario (VENCIDO/HOY/URGENTE)
                nuevo_est = calcular_estado(fecha_limite, fecha_pago or old_fp, old_estado)
                if nuevo_est != old_estado and old_estado not in ("APROBADO", "RECHAZADO"):
                    c.execute(
                        "UPDATE pagos SET estado=?, actualizado_en=CURRENT_TIMESTAMP WHERE id=?",
                        (nuevo_est, pago_id),
                    )
                    actualizados += 1
                continue
            # Si dashboard ya marcó PAGADO y Excel aún no, no borrar fecha_pago local
            if old_estado == "PAGADO" and not fecha_pago and old_fp:
                continue
            # Preservar APROBADO/RECHAZADO si Excel no trae pago
            if old_estado in ("APROBADO", "RECHAZADO") and not fecha_pago:
                estado = old_estado
            c.execute(
                """UPDATE pagos SET
                    prioridad=?, dia_semana=?, fecha_limite=?, detalle=?, proveedor=?,
                    periodo=?, anio=?, centro_costos=?, valor_a_pagar=?, valor_pagado=?,
                    fecha_pago=COALESCE(?, fecha_pago), cuenta_origen_destino=?,
                    pago_total=?, observaciones=?, periodicidad=?, estado=?,
                    sync_hash=?, actualizado_en=CURRENT_TIMESTAMP
                WHERE id=?""",
                (
                    row["prioridad"],
                    row["dia_semana"],
                    row["fecha_limite"],
                    row["detalle"],
                    row["proveedor"],
                    row["periodo"],
                    row["anio"],
                    row["centro_costos"],
                    row["valor_a_pagar"],
                    row["valor_pagado"],
                    row["fecha_pago"],
                    row["cuenta_origen_destino"],
                    row["pago_total"],
                    row["observaciones"],
                    row["periodicidad"],
                    estado,
                    sync_hash,
                    pago_id,
                ),
            )
            actualizados += 1

    # Recalcular estados de todos los no pagados (día cambió)
    c.execute(
        """SELECT id, fecha_limite, fecha_pago, estado FROM pagos
           WHERE estado NOT IN ('PAGADO','RECHAZADO') OR fecha_pago IS NULL"""
    )
    for pid, fl, fp, est in c.fetchall():
        fl_d = _to_date(fl)
        nuevo = calcular_estado(fl_d, fp, est)
        if nuevo != est:
            c.execute(
                "UPDATE pagos SET estado=?, actualizado_en=CURRENT_TIMESTAMP WHERE id=?",
                (nuevo, pid),
            )

    c.execute(
        """INSERT INTO pagos_sync_meta (id, ultima_importacion, registros)
           VALUES (1, CURRENT_TIMESTAMP, ?)
           ON CONFLICT(id) DO UPDATE SET
             ultima_importacion=CURRENT_TIMESTAMP, registros=excluded.registros""",
        (procesados,),
    )
    conn.commit()
    conn.close()
    registrar_accion(
        "pagos",
        "importar_excel_a_sqlite",
        f"procesados={procesados} nuevos={nuevos} actualizados={actualizados}",
        "EXITOSO",
    )
    return {"procesados": procesados, "nuevos": nuevos, "actualizados": actualizados}


def exportar_sqlite_a_excel(path: Path | None = None) -> dict[str, int]:
    """
    Escribe en el Excel valor_pagado (J) y fecha_pago (K) de pagos PAGADO
    que tienen excel_fila, preservando formato con openpyxl.
    """
    path = path or ruta_excel_pagos()
    if not path.is_file():
        raise FileNotFoundError(f"No se encontró Excel para exportar: {path}")

    conn = get_conn()
    c = conn.cursor()
    c.execute(
        """SELECT excel_fila, valor_pagado, fecha_pago, cuenta_origen_destino, pago_total
           FROM pagos
           WHERE estado='PAGADO' AND excel_fila IS NOT NULL AND fecha_pago IS NOT NULL"""
    )
    filas = c.fetchall()
    conn.close()

    if not filas:
        return {"exportados": 0}

    wb = load_workbook(path)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    exportados = 0
    for excel_fila, valor_pagado, fecha_pago, cuenta, pago_total in filas:
        if not excel_fila:
            continue
        # Col J=10 valor pagado, K=11 fecha pago, L=12 cuenta, M=13 pago total
        if valor_pagado is not None:
            ws.cell(row=int(excel_fila), column=10, value=float(valor_pagado))
        fp = _to_date(fecha_pago)
        if fp:
            ws.cell(row=int(excel_fila), column=11, value=datetime(fp.year, fp.month, fp.day))
        if cuenta:
            ws.cell(row=int(excel_fila), column=12, value=cuenta)
        if pago_total:
            ws.cell(row=int(excel_fila), column=13, value=pago_total)
        exportados += 1

    wb.save(path)
    conn = get_conn()
    conn.execute(
        """INSERT INTO pagos_sync_meta (id, ultima_exportacion)
           VALUES (1, CURRENT_TIMESTAMP)
           ON CONFLICT(id) DO UPDATE SET ultima_exportacion=CURRENT_TIMESTAMP"""
    )
    conn.commit()
    conn.close()
    registrar_accion("pagos", "exportar_sqlite_a_excel", f"exportados={exportados}", "EXITOSO")
    return {"exportados": exportados}


def sincronizar_bidireccional(path: Path | None = None) -> dict[str, int]:
    """1) Excel→BD  2) BD→Excel (pagos marcados)."""
    path = path or ruta_excel_pagos()
    imp = importar_excel_a_sqlite(path)
    try:
        exp = exportar_sqlite_a_excel(path)
    except Exception as e:
        logger.warning("Export Excel pagos: %s", e)
        exp = {"exportados": 0}
    return {
        "nuevos": imp.get("nuevos", 0),
        "actualizados": imp.get("actualizados", 0),
        "procesados": imp.get("procesados", 0),
        "exportados": exp.get("exportados", 0),
    }


def actualizar_estados_pagos() -> int:
    """Recalcula estados por fecha (sin tocar Excel)."""
    inicializar_tablas_pagos()
    conn = get_conn()
    c = conn.cursor()
    c.execute("SELECT id, fecha_limite, fecha_pago, estado FROM pagos")
    n = 0
    for pid, fl, fp, est in c.fetchall():
        nuevo = calcular_estado(_to_date(fl), fp, est)
        if nuevo != est:
            c.execute(
                "UPDATE pagos SET estado=?, actualizado_en=CURRENT_TIMESTAMP WHERE id=?",
                (nuevo, pid),
            )
            n += 1
    conn.commit()
    conn.close()
    return n


def programar_tareas_calendario_pagos() -> int:
    """Registra tareas recurrentes del módulo PAGOS e impuestos detectados."""
    from datetime import timedelta

    from utils.calendario_maestro import agregar_tarea

    ahora = datetime.now()
    creadas = 0

    def _existe(titulo: str) -> bool:
        conn = get_conn()
        c = conn.cursor()
        c.execute(
            "SELECT 1 FROM calendario_maestro WHERE titulo=? AND modulo IN ('PAGOS','IMPUESTOS') AND activa=1 LIMIT 1",
            (titulo,),
        )
        ok = c.fetchone() is not None
        conn.close()
        return ok

    semillas = [
        {
            "titulo": "Revisar pagos del día",
            "modulo": "PAGOS",
            "tipo": "AUTOMATICA",
            "prioridad": "ALTA",
            "fecha_inicio": ahora.replace(hour=8, minute=0, second=0, microsecond=0),
            "recurrencia": "DIAS_HABILES",
            "recurrencia_config": {"hora": "08:00"},
            "funcion_callback": "revisar_pagos_pendientes",
        },
        {
            "titulo": "Alerta: pagos de mañana",
            "modulo": "PAGOS",
            "tipo": "ALERTA",
            "prioridad": "MEDIA",
            "fecha_inicio": ahora.replace(hour=16, minute=0, second=0, microsecond=0),
            "recurrencia": "DIAS_HABILES",
            "recurrencia_config": {"hora": "16:00"},
            "funcion_callback": "alerta_pagos_manana",
        },
        {
            "titulo": "Pago nómina semanal EIF",
            "modulo": "PAGOS",
            "tipo": "RECORDATORIO",
            "prioridad": "CRITICA",
            "fecha_inicio": ahora.replace(hour=8, minute=0, second=0, microsecond=0),
            "recurrencia": "SEMANAL",
            "recurrencia_config": {"hora": "08:00", "dias_semana": [4]},
            "funcion_callback": "alerta_nomina_semanal",
        },
    ]
    for t in semillas:
        if not _existe(t["titulo"]):
            agregar_tarea(**t)
            creadas += 1

    # Impuestos futuros → alerta 15 días antes
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        """SELECT id, detalle, proveedor, fecha_limite FROM pagos
           WHERE UPPER(COALESCE(periodicidad,'')) LIKE '%IMPUESTO%'
             AND fecha_limite >= ?
             AND estado != 'PAGADO'
           ORDER BY fecha_limite""",
        (date.today().isoformat(),),
    )
    for pago_id, detalle, proveedor, fl in c.fetchall():
        fl_d = _to_date(fl)
        if not fl_d:
            continue
        titulo = f"IMPUESTO: {detalle} - {proveedor or 'N/A'}"
        if _existe(titulo):
            continue
        inicio = datetime.combine(fl_d - timedelta(days=15), datetime.min.time()).replace(hour=8)
        agregar_tarea(
            titulo=titulo,
            modulo="IMPUESTOS",
            tipo="VENCIMIENTO",
            prioridad="CRITICA",
            fecha_inicio=inicio,
            funcion_callback="alerta_vencimiento_impuesto",
            parametros_callback={"pago_id": pago_id},
        )
        creadas += 1
    conn.close()
    return creadas
