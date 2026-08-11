"""Estandarización de gastos y periodicidad — reglas deterministas (sin IA)."""
from __future__ import annotations

import logging
import os
import re
import unicodedata
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from config import cfg
from database import get_conn, registrar_accion
from modulos.catalogos_gastos import (
    CATEGORIAS_GASTO,
    inicializar_catalogos_gastos,
    mapa_categorias,
    mapa_periodicidades,
)

logger = logging.getLogger(__name__)

SHEET_NAME = "Hoja1"

MAPA_PERIODICIDADES = {
    "MENSUAL": "MENSUAL",
    "ANUAL": "ANUAL",
    "SEMESTRAL": "SEMESTRAL",
    "BIMESTRAL": "BIMESTRAL",
    "ESPORADICO": "ESPORADICO",
    "ESPORÁDICO": "ESPORADICO",
    "IMPUESTO": "IMPUESTO_BIMESTRAL",
    "IMPUESTO MENSUAL": "IMPUESTO_MENSUAL",
    "IMPUESTO BIMESTRAL": "IMPUESTO_BIMESTRAL",
    "IMPUESTO ANUAL": "IMPUESTO_ANUAL",
    "COMPRA": "COMPRA",
    "REUNION": "REUNION",
    "REUNIÒN": "REUNION",
    "REUNIÓN": "REUNION",
    "CITA MEDICA": "REUNION",
    "CITA MÉDICA": "REUNION",
    "CITA_MEDICA": "REUNION",
    "USB": "SEMESTRAL",
    "EDUCACION_USB": "SEMESTRAL",
    "SEMANAL": "SEMANAL",
    "QUINCENAL": "QUINCENAL",
    "": "ESPORADICO",
}

KEYWORDS_SEMANAL = ["SEMANAL", "SEMANA"]
KEYWORDS_QUINCENAL = ["QUINCENAL", "QUINCENA"]

# Reglas de clasificación (orden = prioridad)
REGLAS_CATEGORIA: list[tuple[str, list[str]]] = [
    (
        "IMPUESTOS_DIAN",
        [
            "IVA", "RETENCION", "RENTA", "DECLARACION", "EXOGENA", "DIAN",
            "ANTICIPO IVA", "CERTIFICADO INGRESOS",
        ],
    ),
    ("IMPUESTOS_SHD", ["ICA", "RETEICA", "PREDIAL", "SHD", "INDUSTRIA COMERCIO"]),
    (
        "NOMINA_RRHH",
        [
            "NOMINA", "PRIMA", "VACACION", "CESANTIA", "DOTACION", "LIQUIDACION",
            "PLANILLA", "COMISION", "CRUCE CTAS", "GIRO CTAS", "PAGO NOMINA",
        ],
    ),
    (
        "CREDITOS_BANCARIOS",
        [
            "CUOTA BANCO", "CUOTA VELEZ", "CMR", "TC 6391", "TC VISA", "TC 0670",
            "NUBANK", "LEASING", "CREDITO JACINTA", "CREDITO MOTO", "BANCO DE BOGOTA",
            "BANCOLOMBIA", "DAVIVIENDA", "FALABELLA", "TARJETA",
        ],
    ),
    (
        "SERVICIOS_PUBLICOS",
        [
            "LUZ OF", "AGUA OF", "AGUA SVH", "GAS 40", "CELSIA", "CODENSA",
            "ACUAVALLE", "ETB WIFI", "CLARO HOGAR", "INTERNET", "EPM",
        ],
    ),
    (
        "TELECOMUNICACIONES",
        [
            "MOVISTAR CORPORATIVO", "ETB DEDICADO", "CLARO OF", "SERVIDOR",
            "HOSTING", "DOMINIO", "TIGO UNE",
        ],
    ),
    ("ARRIENDO_INMUEBLE", ["ARRIENDO", "LEASING SVH", "BODEGAJE"]),
    ("ADMON_INMUEBLE", ["ADMON CENTRUM", "ADMON SVH", "SENDEROS", "ADMINISTRACION", "CENTRUM SVH"]),
    ("SALUD_EPS", ["COOMEVA", "EPS", "MEDICINA PREPAGADA"]),
    (
        "SALUD_CITAS",
        [
            "CITA ", "LABORATORIO", "OFTALMOLOGIA", "ODONTOLOGIA", "OTORRINO",
            "CITOLOGIA", "EXAMEN MEDICO",
        ],
    ),
    (
        "EDUCACION_COLEGIOS",
        [
            "PENSION ", "COLEGIO", "DIANA OESE", "ACADEMIA", "DEPORTIVO CALI",
            "RUTA ANDRES", "CAFETERIA", "UNIFORMES",
        ],
    ),
    (
        "EDUCACION_UNIVERSIDAD",
        [
            "USB", "MATRICULA USB", "DERECHO LABORAL", "EPISTEMOLOGIA", "MOODLE",
            "CLUB CONVERSACION", "MATRICULA",
        ],
    ),
    (
        "VEHICULOS",
        ["SOAT", "TECNICOMECANICA", "GASOLINA", "BMW", "PASE MOTO", "BAVARIAN", "CREDITO MOTO"],
    ),
    ("SEGUROS_POLIZAS", ["SEGURO", "POLIZA", "ARL"]),
    (
        "PERSONAL_FAMILIA",
        [
            "AYUDA MAMI", "AHIJADO", "CUMPLEANOS", "ANIVERSARIO", "DIA DEL PADRE",
            "ENTRENADOR", "TOMAS", "DIA FAMILIA",
        ],
    ),
    ("JURIDICO", ["ABOGADO", "MUNARTH", "PREAVISO", "SIC", "AUDIENCIA"]),
    (
        "TECNOLOGIA_SOFTWARE",
        ["BITRIX", "BUFFER", "CAMINOWEB", "WAZZUP", "SOFTWARE", "SUSCRIPCION AXXIS", "AXXIS"],
    ),
    ("PROVEEDORES_EXT", ["MANT PISCINA", "ASEO", "NORALBA", "EMPLEADA", "PAGAR A ", "PAPELERIA"]),
    (
        "COMPRAS_VARIAS",
        ["LLEGADA", "MERCADO", "ALKOSTO", "GUAYOS", "ZAPATOS", "LIBROS", "COMPRA UNIFORMES"],
    ),
    (
        "TAREAS_ADMIN",
        ["ACTUALIZAR", "GRABAR", "REPORTAR", "ENTREGAR", "INGRESAR TODA", "VERIFICAR", "CORRECCION"],
    ),
]


def ruta_excel_estandarizar() -> Path:
    candidatos: list[Path] = []
    env = os.getenv("ESTANDARIZAR_GASTOS", "").strip()
    if env:
        candidatos.append(Path(env))
    cfg_path = getattr(cfg, "ESTANDARIZAR_GASTOS", None)
    if cfg_path:
        candidatos.append(Path(cfg_path))
    base = Path(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    parent = Path(os.path.dirname(base))
    candidatos.extend(
        [
            base / "ESTANDARIZAR_GASTOS.xlsx",
            base / "ESTANDARIZAR GASTOS.xlsx",
            base / "data" / "ESTANDARIZAR_GASTOS.xlsx",
            parent / "Agente ADmin Shaki" / "ESTANDARIZAR GASTOS.xlsx",
            parent / "Agente ADmin Shaki" / "ESTANDARIZAR_GASTOS.xlsx",
        ]
    )
    for p in candidatos:
        if p.is_file():
            return p
    return candidatos[0] if candidatos else base / "ESTANDARIZAR_GASTOS.xlsx"


def limpiar_texto(texto: Any) -> str:
    if texto is None or (isinstance(texto, float) and pd.isna(texto)):
        return ""
    texto = str(texto).strip().upper()
    texto = texto.replace("#U00D1", "Ñ").replace("#U00F3", "Ó")
    texto = texto.replace("#U00C1", "Á").replace("#U00E9", "É")
    texto = texto.replace("#U00ED", "Í").replace("#U00FA", "Ú")
    # Normalizar tildes rotas / encoding raro
    texto = texto.replace("Ò", "Ó").replace("È", "É").replace("À", "Á")
    texto = unicodedata.normalize("NFKC", texto)
    texto = re.sub(r"\s+", " ", texto).strip()
    return texto


def _periodicidad_impuesto_por_proveedor(detalle: str, proveedor: str, default: str) -> str:
    texto = f"{detalle} {proveedor}"
    if any(k in texto for k in ("SHD", "ICA", "RETEICA", "PREDIAL")):
        return "IMPUESTO_BIMESTRAL"
    if any(k in texto for k in ("RENTA", "EXOGENA")):
        return "IMPUESTO_ANUAL"
    if any(k in texto for k in ("RETENCION", "RETE FUENTE", "RETENC")):
        return "IMPUESTO_MENSUAL"
    if "IVA" in texto:
        return "IMPUESTO_BIMESTRAL"
    if "DIAN" in texto:
        return "IMPUESTO_MENSUAL"
    return default


def normalizar_periodicidad(periodicidad_raw: Any, detalle: Any, proveedor: Any = "") -> str:
    per = limpiar_texto(periodicidad_raw)
    det = limpiar_texto(detalle)
    prov = limpiar_texto(proveedor)

    if any(kw in det for kw in KEYWORDS_SEMANAL):
        return "SEMANAL"
    if any(kw in det for kw in KEYWORDS_QUINCENAL):
        return "QUINCENAL"
    if "LLEGADA" in det:
        return "ESPORADICO"
    if per in ("CITA MEDICA", "CITA MÉDICA", "CITA_MEDICA") or "CITA" in det and "MEDIC" in det:
        return "REUNION"
    if per in ("USB", "EDUCACION_USB") or "USB" in det:
        if "MATRICULA" in det or "SEMESTR" in det:
            return "SEMESTRAL"
        return "ANUAL" if "ANUAL" in det else "SEMESTRAL"

    if per == "IMPUESTO":
        return _periodicidad_impuesto_por_proveedor(det, prov, "IMPUESTO_BIMESTRAL")

    return MAPA_PERIODICIDADES.get(per, "ESPORADICO")


def clasificar_categoria(detalle: Any, proveedor: Any) -> tuple[str, float, bool]:
    texto = limpiar_texto(detalle) + " " + limpiar_texto(proveedor)

    # Reglas de negocio especiales
    if "LLEGADA" in texto:
        return "COMPRAS_VARIAS", 0.9, False
    if any(k in texto for k in ("CITA ", "CITA MEDICA", "LABORATORIO", "OFTALMOLOGIA")):
        return "SALUD_CITAS", 0.9, False
    if "USB" in texto:
        return "EDUCACION_UNIVERSIDAD", 0.9, False

    for categoria, keywords in REGLAS_CATEGORIA:
        matches = sum(1 for kw in keywords if kw in texto)
        if matches >= 2:
            return categoria, 0.95, False
        if matches == 1:
            return categoria, 0.75, False

    # Keywords del catálogo (fallback más amplio)
    for codigo, _label, _icono, _color, kws, _mod in CATEGORIAS_GASTO:
        if codigo == "OTRO" or not kws:
            continue
        matches = sum(1 for kw in kws if kw in texto)
        if matches >= 2:
            return codigo, 0.9, False
        if matches == 1:
            return codigo, 0.7, False

    return "OTRO", 0.3, True


def leer_excel_gastos(path: Path | None = None) -> pd.DataFrame:
    path = path or ruta_excel_estandarizar()
    if not path.is_file():
        raise FileNotFoundError(f"No se encontró ESTANDARIZAR GASTOS: {path}")
    df = pd.read_excel(path, sheet_name=SHEET_NAME, header=0)
    cols = {str(c).strip().upper(): c for c in df.columns}
    rename = {}
    for want, aliases in [
        ("detalle", ["DETALLE"]),
        ("proveedor", ["PROVEEDOR"]),
        ("periodicidad", ["PERIODICIDAD"]),
    ]:
        for a in aliases:
            if a in cols:
                rename[cols[a]] = want
                break
    df = df.rename(columns=rename)
    for col in ("detalle", "proveedor", "periodicidad"):
        if col not in df.columns:
            df[col] = None
    return df[["detalle", "proveedor", "periodicidad"]]


def gastos_maestro_vacio() -> bool:
    inicializar_catalogos_gastos()
    conn = get_conn()
    n = conn.execute("SELECT COUNT(*) FROM gastos_maestro").fetchone()[0]
    conn.close()
    return n == 0


def estandarizar_todos(path: Path | None = None, reemplazar: bool = True) -> dict[str, Any]:
    """Lee Excel, limpia/clasifica e inserta en gastos_maestro."""
    inicializar_catalogos_gastos()
    df = leer_excel_gastos(path)

    resultados: dict[str, Any] = {
        "total": 0,
        "clasificados_auto": 0,
        "requieren_revision": 0,
        "periodicidades_corregidas": 0,
        "por_categoria": {},
        "por_periodicidad": {},
    }

    registros: list[tuple] = []
    for _, row in df.iterrows():
        if not str(row.get("detalle") or "").strip() or str(row.get("detalle")).lower() == "nan":
            continue

        det_orig = str(row["detalle"]).strip()
        prov_raw = row.get("proveedor")
        per_raw = row.get("periodicidad")
        prov_orig = "" if prov_raw is None or (isinstance(prov_raw, float) and pd.isna(prov_raw)) else str(prov_raw).strip()
        per_orig = "" if per_raw is None or (isinstance(per_raw, float) and pd.isna(per_raw)) else str(per_raw).strip()

        detalle_limpio = limpiar_texto(det_orig)
        proveedor_limpio = limpiar_texto(prov_orig)
        per_std = normalizar_periodicidad(per_orig, det_orig, prov_orig)
        categoria, confianza, req_revision = clasificar_categoria(det_orig, prov_orig)

        if limpiar_texto(per_orig) != per_std and limpiar_texto(per_orig) not in ("", per_std):
            resultados["periodicidades_corregidas"] += 1
        if req_revision:
            resultados["requieren_revision"] += 1
        else:
            resultados["clasificados_auto"] += 1

        resultados["por_categoria"][categoria] = resultados["por_categoria"].get(categoria, 0) + 1
        resultados["por_periodicidad"][per_std] = resultados["por_periodicidad"].get(per_std, 0) + 1
        resultados["total"] += 1

        registros.append(
            (
                det_orig,
                detalle_limpio,
                prov_orig,
                proveedor_limpio,
                per_orig,
                per_std,
                categoria,
                1,
                int(req_revision),
                float(confianza),
                "EXCEL",
            )
        )

    conn = get_conn()
    if reemplazar:
        conn.execute("DELETE FROM gastos_maestro")
    conn.executemany(
        """INSERT INTO gastos_maestro (
            detalle_original, detalle_limpio, proveedor_original, proveedor_limpio,
            periodicidad_original, periodicidad_std, categoria,
            clasificado_auto, requiere_revision, confianza, origen
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?)""",
        registros,
    )
    conn.commit()
    conn.close()

    registrar_accion(
        "gastos",
        "estandarizar_todos",
        f"total={resultados['total']} revision={resultados['requieren_revision']}",
        "EXITOSO",
    )
    return resultados


def query_gastos(
    categorias: list[str] | None = None,
    periodicidades: list[str] | None = None,
    buscar: str | None = None,
    solo_revision: bool = False,
) -> pd.DataFrame:
    inicializar_catalogos_gastos()
    sql = "SELECT * FROM gastos_maestro WHERE 1=1"
    params: list[Any] = []
    if categorias:
        sql += f" AND categoria IN ({','.join('?' * len(categorias))})"
        params.extend(categorias)
    if periodicidades:
        sql += f" AND periodicidad_std IN ({','.join('?' * len(periodicidades))})"
        params.extend(periodicidades)
    if buscar:
        sql += " AND (detalle_limpio LIKE ? OR proveedor_limpio LIKE ?)"
        params.extend([f"%{buscar.upper()}%", f"%{buscar.upper()}%"])
    if solo_revision:
        sql += " AND requiere_revision=1"
    sql += " ORDER BY requiere_revision DESC, categoria, detalle_limpio"
    conn = get_conn()
    df = pd.read_sql(sql, conn, params=params)
    conn.close()
    return df


def resumen_por_categoria() -> pd.DataFrame:
    conn = get_conn()
    df = pd.read_sql(
        """
        SELECT g.categoria,
               c.label, c.icono, c.color_hex,
               COUNT(*) AS qty,
               GROUP_CONCAT(DISTINCT g.periodicidad_std) AS periodicidades
        FROM gastos_maestro g
        LEFT JOIN cat_categorias_gasto c ON c.codigo = g.categoria
        GROUP BY g.categoria
        ORDER BY qty DESC
        """,
        conn,
    )
    conn.close()
    return df


def resumen_por_periodicidad() -> pd.DataFrame:
    conn = get_conn()
    df = pd.read_sql(
        """
        SELECT g.periodicidad_std AS codigo,
               p.label, p.icono,
               COUNT(*) AS qty
        FROM gastos_maestro g
        LEFT JOIN cat_periodicidades p ON p.codigo = g.periodicidad_std
        GROUP BY g.periodicidad_std
        ORDER BY qty DESC
        """,
        conn,
    )
    conn.close()
    return df


def actualizar_categoria_gasto(gasto_id: int, categoria: str, usuario: str = "ADMIN") -> None:
    conn = get_conn()
    conn.execute(
        """UPDATE gastos_maestro SET
            categoria=?, clasificado_auto=0, requiere_revision=0,
            confianza=1.0, fecha_clasificacion=CURRENT_TIMESTAMP
           WHERE id=?""",
        (categoria, gasto_id),
    )
    conn.commit()
    conn.close()
    registrar_accion("gastos", "actualizar_categoria", f"id={gasto_id} → {categoria}", "EXITOSO")


def marcar_restantes_como_otro() -> int:
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        """UPDATE gastos_maestro SET
            categoria='OTRO', clasificado_auto=0, requiere_revision=0,
            fecha_clasificacion=CURRENT_TIMESTAMP
           WHERE requiere_revision=1"""
    )
    n = c.rowcount
    conn.commit()
    conn.close()
    return n


_LOOKUP_CACHE: dict[str, dict[str, Any] | None] = {}


def lookup_gasto_por_detalle(detalle: str) -> dict[str, Any] | None:
    """Enriquece un pago con categoría/periodicidad estándar (con caché en memoria)."""
    limpio = limpiar_texto(detalle)
    if not limpio:
        return None
    if limpio in _LOOKUP_CACHE:
        return _LOOKUP_CACHE[limpio]
    conn = get_conn()
    row = conn.execute(
        """SELECT categoria, periodicidad_std, requiere_revision
           FROM gastos_maestro WHERE detalle_limpio=? LIMIT 1""",
        (limpio,),
    ).fetchone()
    if not row:
        # match parcial
        row = conn.execute(
            """SELECT categoria, periodicidad_std, requiere_revision
               FROM gastos_maestro
               WHERE detalle_limpio LIKE ? OR ? LIKE '%' || detalle_limpio || '%'
               LIMIT 1""",
            (f"%{limpio[:40]}%", limpio),
        ).fetchone()
    conn.close()
    if not row:
        _LOOKUP_CACHE[limpio] = None
        return None
    cats = mapa_categorias()
    pers = mapa_periodicidades()
    cat = cats.get(row[0], {})
    per = pers.get(row[1], {})
    result = {
        "categoria": row[0],
        "categoria_label": cat.get("label", row[0]),
        "categoria_icono": cat.get("icono", ""),
        "categoria_color": cat.get("color", "#AEAEB2"),
        "periodicidad_std": row[1],
        "periodicidad_label": per.get("label", row[1]),
        "periodicidad_icono": per.get("icono", ""),
        "requiere_revision": bool(row[2]),
    }
    _LOOKUP_CACHE[limpio] = result
    return result


def exportar_excel_estandarizado() -> tuple[bytes, str]:
    df = query_gastos()
    resumen = resumen_por_categoria()
    cats_map = mapa_categorias()

    wb = Workbook()
    ws = wb.active
    ws.title = "Gastos estandarizados"
    headers = [
        "Detalle original",
        "Detalle limpio",
        "Proveedor",
        "Categoría",
        "Periodicidad estándar",
        "Requiere revisión",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1C1C1E")

    for _, r in df.iterrows():
        cat = r.get("categoria") or "OTRO"
        color = (cats_map.get(cat) or {}).get("color", "C7C7CC").lstrip("#")
        ws.append(
            [
                r.get("detalle_original"),
                r.get("detalle_limpio"),
                r.get("proveedor_limpio") or r.get("proveedor_original"),
                cat,
                r.get("periodicidad_std"),
                "SI" if r.get("requiere_revision") else "NO",
            ]
        )
        fill = PatternFill("solid", fgColor=color)
        ws.cell(row=ws.max_row, column=4).fill = fill

    ws2 = wb.create_sheet("Resumen por categoría")
    ws2.append(["Categoría", "Label", "Cantidad", "Periodicidades"])
    for _, r in resumen.iterrows():
        ws2.append([r.get("categoria"), r.get("label"), r.get("qty"), r.get("periodicidades")])

    ws3 = wb.create_sheet("Catálogo periodicidades")
    ws3.append(["Código", "Label", "Icono"])
    for codigo, info in mapa_periodicidades().items():
        ws3.append([codigo, info["label"], info["icono"]])

    bio = BytesIO()
    wb.save(bio)
    nombre = f"GASTOS_ESTANDARIZADOS_{date.today().isoformat()}.xlsx"
    return bio.getvalue(), nombre


def sincronizar_periodicidad_control_pagos() -> dict[str, int]:
    """
    Cruza gastos_maestro con tabla pagos por detalle limpio
    y actualiza periodicidad + intenta escribir en CONTROL_DE_PAGOS.xlsx.
    """
    from utils.sync_pagos import ruta_excel_pagos, SHEET_NAME as PAGOS_SHEET

    inicializar_catalogos_gastos()
    conn = get_conn()
    c = conn.cursor()
    c.execute(
        """SELECT g.detalle_limpio, g.periodicidad_std, g.categoria
           FROM gastos_maestro g"""
    )
    mapa = {row[0]: (row[1], row[2]) for row in c.fetchall()}

    actualizados_bd = 0
    c.execute("SELECT id, detalle, excel_fila FROM pagos")
    pagos = c.fetchall()
    for pago_id, detalle, excel_fila in pagos:
        limpio = limpiar_texto(detalle)
        if limpio not in mapa:
            continue
        per_std, _cat = mapa[limpio]
        c.execute(
            """UPDATE pagos SET periodicidad=?, actualizado_en=CURRENT_TIMESTAMP
               WHERE id=? AND (periodicidad IS NULL OR periodicidad != ?)""",
            (per_std, pago_id, per_std),
        )
        if c.rowcount:
            actualizados_bd += 1
    conn.commit()
    conn.close()

    exportados = 0
    path = ruta_excel_pagos()
    if path.is_file():
        try:
            wb = load_workbook(path)
            ws = wb[PAGOS_SHEET] if PAGOS_SHEET in wb.sheetnames else wb.active
            # Col O = 15 PERIODICIDAD; detalle en D=4
            for row_idx in range(3, ws.max_row + 1):
                det = ws.cell(row=row_idx, column=4).value
                limpio = limpiar_texto(det)
                if limpio in mapa:
                    per_std, _ = mapa[limpio]
                    actual = ws.cell(row=row_idx, column=15).value
                    if limpiar_texto(actual) != per_std:
                        ws.cell(row=row_idx, column=15, value=per_std)
                        exportados += 1
            wb.save(path)
        except Exception as e:
            logger.warning("No se pudo actualizar Excel CONTROL DE PAGOS: %s", e)

    registrar_accion(
        "gastos",
        "sincronizar_periodicidad_control_pagos",
        f"bd={actualizados_bd} excel={exportados}",
        "EXITOSO",
    )
    return {"actualizados_bd": actualizados_bd, "actualizados_excel": exportados}
