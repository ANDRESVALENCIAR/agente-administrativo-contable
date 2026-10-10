"""
Control de pagos — libro "Control de pagos para Agente Shaki".

Fuente: hojas "Maestro de Pagos" (una fila por obligación), "Calendario Anual"
(una fila por cada vez que hay que pagar, llave ID + Ocurrencia) y "Festivos".

Flujo:
1. importar_excel()  → carga/actualiza la tabla `pagos_calendario` (SQLite).
2. enviar_alertas_24h() → 24 h antes de la fecha sugerida de pago envía correo y
   WhatsApp "PENDIENTE DE PAGO" con un botón/enlace "Ya lo pagué".
3. El enlace abre el dashboard (?pagar=<token>) y, al confirmar, marcar_pagado()
   deja la cuota en PAGADO y la escribe de vuelta en el Excel.
"""
from __future__ import annotations

import calendar
import hashlib
import hmac
import logging
import secrets
import shutil
import sqlite3
from datetime import date, datetime, timedelta
from html import escape
from pathlib import Path
from typing import Any

from config import cfg
from database import crear_alerta, get_conn, registrar_accion

logger = logging.getLogger(__name__)

HOJA_MAESTRO = "Maestro de Pagos"
HOJA_CALENDARIO = "Calendario Anual"
HOJA_FESTIVOS = "Festivos"
HOJA_INSTRUCCIONES = "Instrucciones"

ESTADOS = ("PENDIENTE", "EN PROCESO", "PAGADO", "ANULADO")
ABIERTOS = ("PENDIENTE", "EN PROCESO")
MESES = [
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE",
]
DIAS = ["LUNES", "MARTES", "MIÉRCOLES", "JUEVES", "VIERNES", "SÁBADO", "DOMINGO"]

# Columnas del "Calendario Anual" que el agente puede escribir (instrucción 10 del libro)
COL_VALOR, COL_ESTADO, COL_FECHA_PAGO, COL_OBS = "Valor", "Estado", "Fecha pago real", "Observaciones"


# ───────────────────────── Base de datos ─────────────────────────

def inicializar_tablas() -> None:
    conn = get_conn()
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS pagos_calendario (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            id_pago TEXT NOT NULL,
            ocurrencia INTEGER NOT NULL,
            anio INTEGER NOT NULL,
            periodicidad TEXT,
            fecha_limite DATE NOT NULL,
            fecha_sugerida DATE NOT NULL,
            prioridad INTEGER,
            tipo TEXT,
            detalle TEXT,
            proveedor TEXT,
            referencia TEXT,
            centro_costos TEXT,
            valor REAL,
            estado TEXT DEFAULT 'PENDIENTE',
            fecha_pago_real DATE,
            observaciones TEXT,
            excel_fila INTEGER,
            excel_pendiente INTEGER DEFAULT 0,
            confirmado_por TEXT,
            alerta_24h_enviada TIMESTAMP,
            actualizado_en TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            UNIQUE(id_pago, ocurrencia, anio)
        );
        CREATE INDEX IF NOT EXISTS idx_pcal_sugerida ON pagos_calendario(fecha_sugerida);
        CREATE INDEX IF NOT EXISTS idx_pcal_estado ON pagos_calendario(estado);

        CREATE TABLE IF NOT EXISTS pagos_calendario_alertas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            pago_id INTEGER REFERENCES pagos_calendario(id),
            canal TEXT,
            destino TEXT,
            ok INTEGER,
            enviado_en TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS pagos_calendario_meta (
            clave TEXT PRIMARY KEY,
            valor TEXT
        );
        """
    )
    conn.commit()
    conn.close()


def _meta(clave: str, valor: str | None = None) -> str | None:
    conn = get_conn()
    if valor is not None:
        conn.execute(
            "INSERT INTO pagos_calendario_meta(clave, valor) VALUES(?,?) "
            "ON CONFLICT(clave) DO UPDATE SET valor=excluded.valor",
            (clave, valor),
        )
        conn.commit()
    row = conn.execute("SELECT valor FROM pagos_calendario_meta WHERE clave=?", (clave,)).fetchone()
    conn.close()
    return row[0] if row else None


def hay_datos() -> bool:
    inicializar_tablas()
    conn = get_conn()
    n = conn.execute("SELECT COUNT(*) FROM pagos_calendario").fetchone()[0]
    conn.close()
    return n > 0


# ───────────────────────── Fechas (reglas del libro) ─────────────────────────

def fecha_limite(periodicidad: str, ocurrencia: int, anio: int, dia_mes: Any, dia_semana: Any) -> date:
    """Regla 5 (MENSUAL: día del mes, o último día si no existe) y 6 (SEMANAL: n-ésimo día de semana)."""
    if str(periodicidad).upper() == "SEMANAL":
        objetivo = int(dia_semana or 5) - 1  # 1=lunes … 7=domingo
        inicio = date(anio, 1, 1)
        primero = inicio + timedelta(days=(objetivo - inicio.weekday()) % 7)
        return primero + timedelta(weeks=int(ocurrencia) - 1)
    mes = int(ocurrencia)
    ultimo = calendar.monthrange(anio, mes)[1]
    return date(anio, mes, min(int(dia_mes or 1), ultimo))


def fecha_sugerida(limite: date, festivos: set[date]) -> date:
    """Regla 7: si cae sábado, domingo o festivo se adelanta al día hábil anterior."""
    d = limite
    while d.weekday() >= 5 or d in festivos:
        d -= timedelta(days=1)
    return d


# ───────────────────────── Excel ─────────────────────────

def ruta_excel() -> Path:
    return Path(cfg.PAGOS_CONTROL_EXCEL)


def _encabezados(ws) -> dict[str, int]:
    return {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}


def _a_fecha(v: Any) -> date | None:
    if v in (None, ""):
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    try:
        return datetime.fromisoformat(str(v)[:10]).date()
    except ValueError:
        return None


def _a_float(v: Any) -> float | None:
    if v in (None, ""):
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _txt(v: Any) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def leer_libro(path: Path | None = None) -> list[dict[str, Any]]:
    """Lee el libro y devuelve una fila por cuota del calendario, con fechas calculadas."""
    from openpyxl import load_workbook

    path = Path(path or ruta_excel())
    wb = load_workbook(path, data_only=False)
    for hoja in (HOJA_MAESTRO, HOJA_CALENDARIO):
        if hoja not in wb.sheetnames:
            raise ValueError(f"El libro no tiene la hoja '{hoja}'")

    anio = date.today().year
    if HOJA_INSTRUCCIONES in wb.sheetnames:
        v = wb[HOJA_INSTRUCCIONES]["B3"].value
        if isinstance(v, (int, float)) and 2000 < int(v) < 2100:
            anio = int(v)

    festivos: set[date] = set()
    if HOJA_FESTIVOS in wb.sheetnames:
        for (f, *_r) in wb[HOJA_FESTIVOS].iter_rows(min_row=2, values_only=True):
            d = _a_fecha(f)
            if d:
                festivos.add(d)

    wm = wb[HOJA_MAESTRO]
    hm = _encabezados(wm)
    maestro: dict[str, dict[str, Any]] = {}
    for fila in wm.iter_rows(min_row=2, values_only=True):
        r = {k: fila[i - 1] for k, i in hm.items() if i - 1 < len(fila)}
        pid = _txt(r.get("ID"))
        if pid:
            maestro[pid] = r

    wc = wb[HOJA_CALENDARIO]
    hc = _encabezados(wc)
    salida: list[dict[str, Any]] = []
    for n_fila, fila in enumerate(wc.iter_rows(min_row=2, values_only=True), start=2):
        r = {k: fila[i - 1] for k, i in hc.items() if i - 1 < len(fila)}
        pid = _txt(r.get("ID"))
        if not pid or r.get("Ocurrencia") in (None, ""):
            continue
        m = maestro.get(pid)
        if not m:
            logger.warning("Calendario fila %s: ID %s no está en el maestro", n_fila, pid)
            continue
        oc = int(r["Ocurrencia"])
        per = (_txt(m.get("Periodicidad")) or "MENSUAL").upper()
        try:
            lim = fecha_limite(per, oc, anio, m.get("Día del mes"), m.get("Día de la semana"))
        except (ValueError, TypeError) as exc:
            logger.warning("Fila %s (%s-%s): fecha inválida (%s)", n_fila, pid, oc, exc)
            continue
        estado = (_txt(r.get(COL_ESTADO)) or "PENDIENTE").upper()
        salida.append(
            {
                "id_pago": pid,
                "ocurrencia": oc,
                "anio": anio,
                "periodicidad": per,
                "fecha_limite": lim,
                "fecha_sugerida": fecha_sugerida(lim, festivos),
                "prioridad": int(m.get("Prioridad") or 3),
                "tipo": (_txt(m.get("Tipo")) or "PAGO").upper(),
                "detalle": _txt(m.get("Detalle")),
                "proveedor": _txt(m.get("Proveedor")),
                "referencia": _txt(m.get("Referencia / Cuenta")),
                "centro_costos": _txt(m.get("Centro de costos")),
                "valor": _a_float(r.get(COL_VALOR)),
                "estado": estado if estado in ESTADOS else "PENDIENTE",
                "fecha_pago_real": _a_fecha(r.get(COL_FECHA_PAGO)),
                "observaciones": _txt(r.get(COL_OBS)),
                "excel_fila": n_fila,
            }
        )
    return salida


def importar_excel(path: Path | None = None) -> dict[str, int]:
    """
    Carga el libro en SQLite. Idempotente (llave ID + Ocurrencia + año).
    Si en la BD una cuota ya quedó PAGADA pero el Excel aún no lo refleja
    (escritura pendiente), se respeta lo de la BD.
    """
    inicializar_tablas()
    path = Path(path or ruta_excel())
    filas = leer_libro(path)
    conn = get_conn()
    c = conn.cursor()
    nuevos = actualizados = 0
    for f in filas:
        prev = c.execute(
            "SELECT id, estado, excel_pendiente FROM pagos_calendario WHERE id_pago=? AND ocurrencia=? AND anio=?",
            (f["id_pago"], f["ocurrencia"], f["anio"]),
        ).fetchone()
        valores = (
            f["periodicidad"], f["fecha_limite"].isoformat(), f["fecha_sugerida"].isoformat(), f["prioridad"],
            f["tipo"], f["detalle"], f["proveedor"], f["referencia"], f["centro_costos"], f["excel_fila"],
        )
        if prev is None:
            c.execute(
                """INSERT INTO pagos_calendario
                   (periodicidad, fecha_limite, fecha_sugerida, prioridad, tipo, detalle, proveedor,
                    referencia, centro_costos, excel_fila, id_pago, ocurrencia, anio,
                    valor, estado, fecha_pago_real, observaciones)
                   VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                valores + (
                    f["id_pago"], f["ocurrencia"], f["anio"], f["valor"], f["estado"],
                    f["fecha_pago_real"].isoformat() if f["fecha_pago_real"] else None, f["observaciones"],
                ),
            )
            nuevos += 1
            continue
        c.execute(
            """UPDATE pagos_calendario SET periodicidad=?, fecha_limite=?, fecha_sugerida=?, prioridad=?,
                   tipo=?, detalle=?, proveedor=?, referencia=?, centro_costos=?, excel_fila=?,
                   actualizado_en=CURRENT_TIMESTAMP WHERE id=?""",
            valores + (prev[0],),
        )
        if not prev[2]:  # sin escritura pendiente: el Excel manda en los campos editables
            c.execute(
                """UPDATE pagos_calendario SET valor=?, estado=?, fecha_pago_real=?, observaciones=?
                   WHERE id=?""",
                (
                    f["valor"], f["estado"],
                    f["fecha_pago_real"].isoformat() if f["fecha_pago_real"] else None,
                    f["observaciones"], prev[0],
                ),
            )
        actualizados += 1
    conn.commit()
    conn.close()
    _meta("ultima_importacion", datetime.now().isoformat(timespec="seconds"))
    _meta("ruta_excel", str(path))
    registrar_accion("pagos", "importar_control_pagos", f"{nuevos} nuevos, {actualizados} actualizados", "EXITOSO")
    return {"nuevos": nuevos, "actualizados": actualizados, "total": len(filas)}


def importar_si_existe(forzar: bool = False) -> dict[str, int] | None:
    """Importa el libro configurado si existe (tarea programada / arranque).
    Si el Excel no ha cambiado desde la última lectura no se vuelve a leer (salvo forzar=True)."""
    p = ruta_excel()
    if not p.is_file():
        logger.info("Control de pagos: no se encontró %s", p)
        return None
    st = p.stat()
    huella = f"{p}|{st.st_mtime_ns}|{st.st_size}"
    if not forzar and _meta("huella_excel") == huella:
        escribir_pendientes_excel()
        return {"nuevos": 0, "actualizados": 0, "total": 0, "sin_cambios": True}
    res = importar_excel(p)
    _meta("huella_excel", huella)
    escribir_pendientes_excel()
    return res


def _escribir_fila_excel(path: Path, fila: int, estado: str, fecha_pago: str | None,
                         observaciones: str | None, valor: float | None) -> None:
    from openpyxl import load_workbook

    respaldo = path.with_suffix(path.suffix + ".bak")
    if not respaldo.exists():
        shutil.copy2(path, respaldo)  # copia intacta antes de la primera escritura del agente
    wb = load_workbook(path)
    ws = wb[HOJA_CALENDARIO]
    h = _encabezados(ws)
    ws.cell(row=fila, column=h[COL_ESTADO], value=estado)
    if COL_FECHA_PAGO in h:
        ws.cell(row=fila, column=h[COL_FECHA_PAGO],
                value=datetime.fromisoformat(fecha_pago) if fecha_pago else None)
        if fecha_pago:
            ws.cell(row=fila, column=h[COL_FECHA_PAGO]).number_format = "dd/mm/yyyy"
    if COL_OBS in h and observaciones is not None:
        ws.cell(row=fila, column=h[COL_OBS], value=observaciones)
    if COL_VALOR in h and valor is not None:
        ws.cell(row=fila, column=h[COL_VALOR], value=valor)
    wb.save(path)


def escribir_pendientes_excel() -> int:
    """Reintenta escribir en el Excel las confirmaciones que no se pudieron guardar (archivo abierto)."""
    if not cfg.PAGOS_ESCRIBIR_EXCEL:
        return 0
    path = ruta_excel()
    if not path.is_file():
        return 0
    conn = get_conn()
    rows = conn.execute(
        """SELECT id, excel_fila, estado, fecha_pago_real, observaciones, valor
           FROM pagos_calendario WHERE excel_pendiente=1 AND excel_fila IS NOT NULL"""
    ).fetchall()
    conn.close()
    ok = 0
    for pid, fila, estado, fpago, obs, valor in rows:
        try:
            _escribir_fila_excel(path, fila, estado, fpago, obs, valor)
        except PermissionError:
            logger.warning("Control de pagos: Excel abierto, se reintentará (%s)", path)
            break
        except Exception as exc:
            logger.error("Control de pagos: no se pudo escribir fila %s: %s", fila, exc)
            continue
        conn = get_conn()
        conn.execute("UPDATE pagos_calendario SET excel_pendiente=0 WHERE id=?", (pid,))
        conn.commit()
        conn.close()
        ok += 1
    return ok


# ───────────────────────── Consultas y estados ─────────────────────────

COLUMNAS = [
    "id", "id_pago", "ocurrencia", "anio", "periodicidad", "fecha_limite", "fecha_sugerida", "prioridad",
    "tipo", "detalle", "proveedor", "referencia", "centro_costos", "valor", "estado", "fecha_pago_real",
    "observaciones", "excel_fila", "excel_pendiente", "confirmado_por", "alerta_24h_enviada",
]


def obtener(pago_id: int) -> dict[str, Any] | None:
    conn = get_conn()
    row = conn.execute(f"SELECT {', '.join(COLUMNAS)} FROM pagos_calendario WHERE id=?", (pago_id,)).fetchone()
    conn.close()
    return dict(zip(COLUMNAS, row)) if row else None


def listar(desde: date | None = None, hasta: date | None = None, estados: tuple[str, ...] | None = None,
           centro: str | None = None) -> list[dict[str, Any]]:
    sql = f"SELECT {', '.join(COLUMNAS)} FROM pagos_calendario WHERE 1=1"
    params: list[Any] = []
    if desde:
        sql += " AND fecha_sugerida >= ?"
        params.append(desde.isoformat())
    if hasta:
        sql += " AND fecha_sugerida <= ?"
        params.append(hasta.isoformat())
    if estados:
        sql += f" AND estado IN ({','.join('?' * len(estados))})"
        params.extend(estados)
    if centro:
        sql += " AND centro_costos = ?"
        params.append(centro)
    sql += " ORDER BY fecha_sugerida, prioridad, id_pago"
    conn = get_conn()
    rows = conn.execute(sql, params).fetchall()
    conn.close()
    return [dict(zip(COLUMNAS, r)) for r in rows]


def semaforo(p: dict[str, Any], hoy: date | None = None) -> str:
    """VENCIDO · HOY · MAÑANA · PRÓXIMO (≤3 días) · AL DÍA · PAGADO · ANULADO."""
    if p["estado"] in ("PAGADO", "ANULADO"):
        return p["estado"]
    hoy = hoy or date.today()
    d = (date.fromisoformat(str(p["fecha_sugerida"])) - hoy).days
    if d < 0:
        return "VENCIDO"
    if d == 0:
        return "HOY"
    if d == 1:
        return "MAÑANA"
    if d <= 3:
        return "PRÓXIMO"
    return "AL DÍA"


def resumen(hoy: date | None = None) -> dict[str, Any]:
    """Números para el dashboard."""
    hoy = hoy or date.today()
    inicializar_tablas()
    conn = get_conn()
    q = lambda sql, *p: conn.execute(sql, p).fetchone()[0]  # noqa: E731
    ab = "estado IN ('PENDIENTE','EN PROCESO')"
    ini_mes = hoy.replace(day=1).isoformat()
    fin_mes = date(hoy.year, hoy.month, calendar.monthrange(hoy.year, hoy.month)[1]).isoformat()
    r = {
        "vencidos": q(f"SELECT COUNT(*) FROM pagos_calendario WHERE {ab} AND fecha_sugerida < ?", hoy.isoformat()),
        "hoy": q(f"SELECT COUNT(*) FROM pagos_calendario WHERE {ab} AND fecha_sugerida = ?", hoy.isoformat()),
        "manana": q(f"SELECT COUNT(*) FROM pagos_calendario WHERE {ab} AND fecha_sugerida = ?",
                    (hoy + timedelta(days=1)).isoformat()),
        "semana": q(f"SELECT COUNT(*) FROM pagos_calendario WHERE {ab} AND fecha_sugerida BETWEEN ? AND ?",
                    hoy.isoformat(), (hoy + timedelta(days=7)).isoformat()),
        "mes_total": q("SELECT COUNT(*) FROM pagos_calendario WHERE fecha_sugerida BETWEEN ? AND ? AND estado!='ANULADO'",
                       ini_mes, fin_mes),
        "mes_pagados": q("SELECT COUNT(*) FROM pagos_calendario WHERE fecha_sugerida BETWEEN ? AND ? AND estado='PAGADO'",
                         ini_mes, fin_mes),
        "mes_pendientes": q(f"SELECT COUNT(*) FROM pagos_calendario WHERE {ab} AND fecha_sugerida BETWEEN ? AND ?",
                            ini_mes, fin_mes),
        "valor_pendiente_mes": q(
            f"SELECT COALESCE(SUM(valor),0) FROM pagos_calendario WHERE {ab} AND fecha_sugerida BETWEEN ? AND ?",
            ini_mes, fin_mes),
        "alertas_hoy": q("SELECT COUNT(*) FROM pagos_calendario_alertas WHERE DATE(enviado_en)=DATE('now','localtime')"),
        "excel_pendiente": q("SELECT COUNT(*) FROM pagos_calendario WHERE excel_pendiente=1"),
    }
    conn.close()
    r["al_dia_pct"] = round(100 * r["mes_pagados"] / r["mes_total"]) if r["mes_total"] else 100
    return r


# ───────────────────────── Confirmación ─────────────────────────

def marcar_pagado(pago_id: int, por: str = "dashboard", valor: float | None = None,
                  referencia: str | None = None, fecha_pago: date | None = None) -> dict[str, Any]:
    """Deja la cuota PAGADA en la BD y en el Excel. Devuelve {'ok', 'ya_estaba', 'excel'}."""
    p = obtener(pago_id)
    if not p:
        return {"ok": False, "motivo": "no_existe"}
    if p["estado"] == "PAGADO":
        return {"ok": True, "ya_estaba": True, "excel": not p["excel_pendiente"]}
    fpago = (fecha_pago or date.today()).isoformat()
    nota = f"Pagado · confirmado por {por} el {datetime.now():%d/%m/%Y %H:%M}"
    if referencia:
        nota += f" · ref {referencia}"
    obs = f"{p['observaciones']} | {nota}" if p.get("observaciones") else nota
    conn = get_conn()
    conn.execute(
        """UPDATE pagos_calendario SET estado='PAGADO', fecha_pago_real=?, observaciones=?,
               valor=COALESCE(?, valor), confirmado_por=?, excel_pendiente=1,
               actualizado_en=CURRENT_TIMESTAMP WHERE id=?""",
        (fpago, obs, valor, por, pago_id),
    )
    conn.commit()
    conn.close()
    registrar_accion("pagos", "marcar_pagado", f"{p['id_pago']}-{p['ocurrencia']} {p['detalle']} ({por})", "EXITOSO")
    escrito = escribir_pendientes_excel() > 0 if cfg.PAGOS_ESCRIBIR_EXCEL else False
    return {"ok": True, "ya_estaba": False, "excel": escrito}


def cambiar_estado(pago_id: int, estado: str, por: str = "dashboard") -> None:
    if estado == "PAGADO":
        marcar_pagado(pago_id, por=por)
        return
    estado = estado.upper()
    if estado not in ESTADOS:
        raise ValueError(estado)
    conn = get_conn()
    conn.execute(
        """UPDATE pagos_calendario SET estado=?, fecha_pago_real=NULL, excel_pendiente=1,
               actualizado_en=CURRENT_TIMESTAMP WHERE id=?""",
        (estado, pago_id),
    )
    conn.commit()
    conn.close()
    escribir_pendientes_excel()


# ───────────────────────── Enlaces firmados ─────────────────────────

def _secreto() -> bytes:
    secreto = getattr(cfg, "PAGOS_LINK_SECRET", None)
    if secreto:
        return secreto.encode()
    inicializar_tablas()
    s = _meta("link_secret")
    if not s:
        s = _meta("link_secret", secrets.token_urlsafe(32))
    return str(s).encode()


def generar_token(pago_id: int) -> str:
    firma = hmac.new(_secreto(), f"pagar:{pago_id}".encode(), hashlib.sha256).hexdigest()[:32]
    return f"{pago_id}.{firma}"


def verificar_token(token: str) -> int | None:
    try:
        pid_txt, firma = str(token).split(".", 1)
        pid = int(pid_txt)
    except (ValueError, AttributeError):
        return None
    esperado = generar_token(pid).split(".", 1)[1]
    return pid if hmac.compare_digest(esperado, firma) else None


def enlace_confirmacion(pago_id: int) -> str:
    return f"{getattr(cfg, 'APP_PUBLIC_URL', 'http://localhost:8501')}/?pagar={generar_token(pago_id)}"


# ───────────────────────── Alertas 24 h ─────────────────────────

def _fecha_humana(d: date) -> str:
    return f"{DIAS[d.weekday()].capitalize()} {d.day} de {MESES[d.month - 1].lower()}"


def _miles(v: float | None) -> str:
    return f"${v:,.0f}".replace(",", ".") if v else "Por definir"


def _cuando(p: dict[str, Any], hoy: date | None = None) -> str:
    d = (date.fromisoformat(str(p["fecha_sugerida"])) - (hoy or date.today())).days
    return "VENCE HOY" if d <= 0 else "VENCE MAÑANA" if d == 1 else f"VENCE EN {d} DÍAS"


def html_correo(pagos: list[dict[str, Any]], hoy: date | None = None) -> str:
    """Correo compatible con Gmail/Outlook: tablas y estilos en línea, un botón por pago."""
    bloques = []
    for p in pagos:
        f = date.fromisoformat(str(p["fecha_sugerida"]))
        accion = "Ya lo pagué" if p["tipo"] == "PAGO" else "Ya lo hice"
        ref = f"<br>Referencia: <b>{escape(p['referencia'])}</b>" if p.get("referencia") else ""
        bloques.append(f"""
        <tr><td style="padding:0 0 16px 0">
          <table role="presentation" width="100%" cellpadding="0" cellspacing="0"
                 style="border:1px solid #E4E4E7;border-radius:12px;background:#FFFFFF" bgcolor="#FFFFFF">
            <tr><td style="padding:18px 20px;font-family:Arial,Helvetica,sans-serif;color:#18181B">
              <div style="font-size:11px;letter-spacing:.08em;color:#C2410C;font-weight:bold">PENDIENTE DE PAGO · {_cuando(p, hoy)}</div>
              <div style="font-size:18px;font-weight:bold;margin:6px 0 4px 0">{escape(p['detalle'] or '')}</div>
              <div style="font-size:14px;color:#52525B;line-height:1.5">
                {escape(p['proveedor'] or '')} · {escape(p['centro_costos'] or '')}{ref}<br>
                Fecha de pago: <b>{_fecha_humana(f)}</b> · Valor: <b>{_miles(p.get('valor'))}</b>
              </div>
              <table role="presentation" cellpadding="0" cellspacing="0" style="margin-top:14px"><tr>
                <td style="background:#18181B;border-radius:999px" bgcolor="#18181B">
                  <a href="{escape(enlace_confirmacion(p['id']))}"
                     style="display:inline-block;padding:12px 22px;font-family:Arial,Helvetica,sans-serif;
                            font-size:14px;font-weight:bold;color:#FFFFFF;text-decoration:none">✓ {accion}</a>
                </td></tr></table>
            </td></tr>
          </table>
        </td></tr>""")
    return f"""<!doctype html><html><body style="margin:0;background:#F4F4F5">
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#F4F4F5">
      <tr><td align="center" style="padding:24px 12px">
        <table role="presentation" width="560" cellpadding="0" cellspacing="0" style="max-width:560px;width:100%">
          <tr><td style="padding:0 0 18px 0;font-family:Arial,Helvetica,sans-serif">
            <div style="font-size:20px;font-weight:bold;color:#18181B">Shaki · {escape(cfg.NOMBRE_EMPRESA)}</div>
            <div style="font-size:14px;color:#52525B;margin-top:4px">
              Tienes {len(pagos)} pago(s) pendiente(s) de pago. Cuando lo pagues, toca el botón y el agente
              lo deja al día.</div>
          </td></tr>
          {''.join(bloques)}
          <tr><td style="font-family:Arial,Helvetica,sans-serif;font-size:12px;color:#A1A1AA;padding-top:6px">
            Mensaje automático del agente administrativo. El botón abre el panel de Shaki para confirmar.
          </td></tr>
        </table>
      </td></tr>
    </table></body></html>"""


def texto_whatsapp(p: dict[str, Any], hoy: date | None = None) -> str:
    f = date.fromisoformat(str(p["fecha_sugerida"]))
    accion = "Ya lo pagué" if p["tipo"] == "PAGO" else "Ya lo hice"
    lineas = [
        f"⏰ *PENDIENTE DE PAGO · {_cuando(p, hoy).lower()}*",
        f"*{p['detalle']}*",
        f"{p['proveedor'] or ''} · {p['centro_costos'] or ''}".strip(" ·"),
    ]
    if p.get("referencia"):
        lineas.append(f"Ref: {p['referencia']}")
    lineas += [
        f"Fecha: {_fecha_humana(f)} · Valor: {_miles(p.get('valor'))}",
        "",
        f"✅ {accion}: {enlace_confirmacion(p['id'])}",
    ]
    return "\n".join(lineas)


def variables_plantilla(p: dict[str, Any]) -> dict[str, str]:
    """Variables para la plantilla aprobada de WhatsApp: {{1}} detalle {{2}} fecha {{3}} valor {{4}} enlace."""
    f = date.fromisoformat(str(p["fecha_sugerida"]))
    return {"1": str(p["detalle"] or ""), "2": _fecha_humana(f), "3": _miles(p.get("valor")),
            "4": enlace_confirmacion(p["id"])}


def texto_correo_plano(pagos: list[dict[str, Any]], hoy: date | None = None) -> str:
    return "\n\n".join(texto_whatsapp(p, hoy).replace("*", "") for p in pagos)


def _registrar_envio(pago_id: int, canal: str, destino: str, ok: bool) -> None:
    conn = get_conn()
    conn.execute(
        "INSERT INTO pagos_calendario_alertas(pago_id, canal, destino, ok) VALUES (?,?,?,?)",
        (pago_id, canal, destino, int(ok)),
    )
    conn.commit()
    conn.close()


def _destinos_correo() -> list[str]:
    return [e.strip() for e in (cfg.PAGOS_ALERTA_EMAILS or "").replace(";", ",").split(",") if e.strip()]


def _destinos_whatsapp() -> list[str]:
    from conexiones.whatsapp_client import _parse_destinatarios

    return _parse_destinatarios(cfg.PAGOS_ALERTA_WHATSAPP)


def estado_canales() -> dict[str, Any]:
    """Qué falta para que las alertas salgan solas (para el panel)."""
    from conexiones.notificador import canal_correo, canal_whatsapp

    url = cfg.APP_PUBLIC_URL
    return {
        "correo": canal_correo(),
        "whatsapp": canal_whatsapp(),
        "destinos_correo": _destinos_correo(),
        "destinos_whatsapp": [d.replace("whatsapp:", "") for d in _destinos_whatsapp()],
        "url_publica": not any(x in url for x in ("localhost", "127.0.0.1", "0.0.0.0")),
        "url": url,
    }


def pagos_para_alerta(hoy: date | None = None) -> list[dict[str, Any]]:
    """
    Cuotas abiertas sin aviso cuya fecha sugerida es mañana (alerta 24 h) u hoy
    (recuperación: si el PC estuvo apagado ayer, se avisa hoy temprano).
    """
    hoy = hoy or date.today()
    conn = get_conn()
    rows = conn.execute(
        f"""SELECT {', '.join(COLUMNAS)} FROM pagos_calendario
            WHERE estado IN ('PENDIENTE','EN PROCESO') AND fecha_sugerida IN (?, ?)
              AND alerta_24h_enviada IS NULL
            ORDER BY fecha_sugerida, prioridad, id_pago""",
        (hoy.isoformat(), (hoy + timedelta(days=1)).isoformat()),
    ).fetchall()
    conn.close()
    return [dict(zip(COLUMNAS, r)) for r in rows]


def _reclamar(ids: list[int]) -> list[int]:
    """Marca como avisadas solo las que nadie más reclamó (evita duplicados entre procesos)."""
    conn = get_conn()
    ok = []
    for i in ids:
        cur = conn.execute(
            "UPDATE pagos_calendario SET alerta_24h_enviada=datetime('now','localtime') "
            "WHERE id=? AND alerta_24h_enviada IS NULL",
            (i,),
        )
        if cur.rowcount == 1:
            ok.append(i)
    conn.commit()
    conn.close()
    return ok


def _liberar(ids: list[int]) -> None:
    conn = get_conn()
    conn.executemany("UPDATE pagos_calendario SET alerta_24h_enviada=NULL WHERE id=?", [(i,) for i in ids])
    conn.commit()
    conn.close()


def enviar_alerta(pagos: list[dict[str, Any]], marcar: bool = True, hoy: date | None = None) -> dict[str, Any]:
    """
    Envía un correo (con todos los pagos y un botón por pago) y un WhatsApp por pago.
    Con marcar=True la cuota queda como avisada SOLO si salió por algún canal real;
    si no hay canal configurado o todo falla, se libera para reintentar en la próxima ronda.
    """
    from conexiones.notificador import enviar_email, enviar_whatsapp

    vacio = {"pagos": 0, "correos": 0, "whatsapp": 0, "sin_canal": False}
    if not pagos:
        return vacio
    if marcar:
        reclamados = set(_reclamar([p["id"] for p in pagos]))
        pagos = [p for p in pagos if p["id"] in reclamados]
        if not pagos:
            return vacio

    correos = whats = 0
    exito: set[int] = set()
    sin_canal = True
    asunto = (
        f"Pendiente de pago: {pagos[0]['detalle']} · {_cuando(pagos[0], hoy).lower()}"
        if len(pagos) == 1 else f"{len(pagos)} pagos pendientes de pago"
    )
    html, plano = html_correo(pagos, hoy), texto_correo_plano(pagos, hoy)
    for dest in _destinos_correo():
        r = enviar_email(dest, asunto, html, plano)
        if r is None:
            continue
        sin_canal = False
        correos += bool(r)
        for p in pagos:
            _registrar_envio(p["id"], "correo", dest, bool(r))
            if r:
                exito.add(p["id"])
    for p in pagos:
        for dest in _destinos_whatsapp():
            r = enviar_whatsapp(dest, texto_whatsapp(p, hoy), variables_plantilla(p))
            if r is None:
                continue
            sin_canal = False
            whats += bool(r)
            _registrar_envio(p["id"], "whatsapp", dest.replace("whatsapp:", ""), bool(r))
            if r:
                exito.add(p["id"])
    if marcar:
        fallidos = [p["id"] for p in pagos if p["id"] not in exito]
        if fallidos:
            _liberar(fallidos)
    return {"pagos": len(exito), "correos": correos, "whatsapp": whats, "sin_canal": sin_canal,
            "pendientes": len(pagos) - len(exito)}


def enviar_alertas_24h(hoy: date | None = None) -> dict[str, Any]:
    """Tarea programada: alerta 24 h antes de la fecha sugerida (y recuperación del mismo día)."""
    inicializar_tablas()
    try:
        importar_si_existe()
    except Exception as exc:  # seguir con lo que hay en la BD
        logger.warning("Control de pagos: no se pudo releer el Excel: %s", exc)
    pagos = pagos_para_alerta(hoy)
    res = enviar_alerta(pagos, hoy=hoy)
    if pagos and res["sin_canal"]:
        crear_alerta(
            "URGENTE", "pagos", "Alertas de pago SIN ENVIAR: falta configurar correo o WhatsApp",
            f"{len(pagos)} cuota(s) por avisar. Configure SMTP_* o Twilio en .env.",
        )
    elif res["pagos"]:
        crear_alerta(
            "AVISO", "pagos", f"Alerta enviada: {res['pagos']} pago(s) por vencer",
            "; ".join(p["detalle"] or "" for p in pagos)[:400],
        )
    registrar_accion("pagos", "alertas_24h", str(res), "EXITOSO" if not res.get("pendientes") else "ERROR")
    return res


def inicializar_y_cargar() -> None:
    """Arranque: tablas + import del Excel si existe (silencioso)."""
    inicializar_tablas()
    try:
        importar_si_existe()
    except (sqlite3.Error, ValueError, OSError) as exc:
        logger.warning("Control de pagos al arrancar: %s", exc)
