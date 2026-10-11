"""
Lectores de insumos del contador para la revisión de impuestos.

- leer_pdf_formulario(): toma el borrador/formulario DIAN en PDF y devuelve {renglón: valor}.
  Los formularios DIAN (350, 300, 110…) imprimen el número de renglón en una casilla y el valor
  a su derecha en la misma línea: se asocia cada cifra con el renglón más cercano a su izquierda.
- leer_excel_formulario(): busca en la liquidación del contador la hoja que replica el
  formulario (celda con el número de renglón y el valor en la celda de al lado).
- Libro: acceso rápido (solo lectura) a las hojas de la liquidación: Prueba, Auxiliar,
  RetenciónHT, AutorrentaHT, Formato… (plantilla mensual del contador de EIF).
"""
from __future__ import annotations

import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

_NUM = re.compile(r"^-?\(?\$?\s*\d{1,3}(?:[.,]\d{3})+\)?$|^-?\d{4,}$|^0$")
_LABEL = re.compile(r"^\d{1,3}$")


def _a_numero(txt: str) -> float | None:
    t = txt.strip().replace("$", "").replace(" ", "")
    neg = t.startswith("-") or (t.startswith("(") and t.endswith(")"))
    t = t.strip("-()")
    if re.fullmatch(r"\d{1,3}(?:[.,]\d{3})+", t):
        t = re.sub(r"[.,]", "", t)
    try:
        v = float(t)
    except ValueError:
        return None
    return -v if neg else v


def norm(txt: Any) -> str:
    """Minúsculas sin tildes ni espacios dobles (para buscar rótulos)."""
    s = unicodedata.normalize("NFKD", str(txt or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


# ───────────────────────── PDF ─────────────────────────

def leer_pdf_formulario(ruta: str | Path) -> dict[str, Any]:
    """Devuelve {"campos": {int: float}, "texto": str, "paginas": n}. Nunca lanza por formato."""
    import pdfplumber

    campos: dict[int, float] = {}
    textos: list[str] = []
    with pdfplumber.open(str(ruta)) as pdf:
        for page in pdf.pages:
            palabras = page.extract_words(keep_blank_chars=False, use_text_flow=False)
            textos.append(page.extract_text() or "")
            etiquetas = [w for w in palabras if _LABEL.match(w["text"])]
            for w in palabras:
                if _LABEL.match(w["text"]) and w["text"] != "0":
                    continue
                if not _NUM.match(w["text"]):
                    continue
                valor = _a_numero(w["text"])
                if valor is None:
                    continue
                cerca = [e for e in etiquetas
                         if abs(e["top"] - w["top"]) <= 6 and e["x1"] <= w["x0"] + 1 and e is not w]
                if not cerca:
                    continue
                e = max(cerca, key=lambda x: x["x0"])
                n = int(e["text"])
                if n not in campos:
                    campos[n] = valor
    return {"campos": campos, "texto": "\n".join(textos), "paginas": len(textos)}


# ───────────────────────── Excel ─────────────────────────

def _es_num(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _filas_hoja(ws) -> list[tuple]:
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


def campos_desde_filas(filas: list[tuple], minimo: int = 26) -> dict[int, float]:
    """Celda con número de renglón entero seguida (columna contigua) de un valor numérico."""
    out: dict[int, float] = {}
    for fila in filas:
        for j in range(len(fila) - 1):
            v, sig = fila[j], fila[j + 1]
            if _es_num(v) and float(v).is_integer() and minimo <= v <= 999 and _es_num(sig):
                out.setdefault(int(v), float(sig))
    return out


def leer_excel_formulario(ruta: str | Path, formulario: str = "") -> dict[str, Any]:
    """Busca la hoja que replica el formulario y devuelve {"hoja", "campos"}."""
    import pandas as pd

    ruta = Path(ruta)
    hojas: dict[str, list[tuple]] = {}
    if ruta.suffix.lower() in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        nombres = wb.sheetnames
        preferidas = [n for n in nombres if re.search(rf"form|{formulario or 'xx'}|declar|borrador", norm(n))]
        for n in preferidas or nombres:
            ws = wb[n]
            if (ws.max_row or 0) > 400:   # hojas de movimientos: no son el formulario
                continue
            hojas[n] = _filas_hoja(ws)
        wb.close()
    else:
        for n, df in pd.read_excel(ruta, sheet_name=None, header=None).items():
            if len(df) <= 400:
                hojas[n] = [tuple(None if pd.isna(x) else x for x in r) for r in df.itertuples(index=False)]
    mejor, campos = None, {}
    for n, filas in hojas.items():
        c = campos_desde_filas(filas)
        if len(c) > len(campos):
            mejor, campos = n, c
    return {"hoja": mejor, "campos": campos}


class Libro:
    """Liquidación mensual del contador (.xlsm/.xlsx), solo lectura y con caché por hoja."""

    def __init__(self, ruta: str | Path):
        import openpyxl

        self.ruta = Path(ruta)
        self.wb = openpyxl.load_workbook(self.ruta, data_only=True, read_only=True)
        self._cache: dict[str, list[tuple]] = {}

    @property
    def hojas(self) -> list[str]:
        return list(self.wb.sheetnames)

    def hoja(self, *nombres: str) -> list[tuple] | None:
        objetivo = [norm(n) for n in nombres]
        for real in self.wb.sheetnames:
            if norm(real) in objetivo:
                if real not in self._cache:
                    self._cache[real] = _filas_hoja(self.wb[real])
                return self._cache[real]
        return None

    def hojas_que_empiezan(self, prefijo: str) -> dict[str, list[tuple]]:
        out = {}
        for real in self.wb.sheetnames:
            if norm(real).startswith(norm(prefijo)):
                if real not in self._cache:
                    self._cache[real] = _filas_hoja(self.wb[real])
                out[real] = self._cache[real]
        return out

    def hoja_con(self, requeridas: tuple[str, ...], filas_max: int = 40) -> list[tuple] | None:
        """Primera hoja que tenga una fila de encabezados con todas las palabras `requeridas` (sirve para
        archivos exportados de World Office u otro programa, donde la hoja se llama distinto)."""
        req = [norm(r) for r in requeridas]
        for real in self.wb.sheetnames:
            if real not in self._cache:
                self._cache[real] = _filas_hoja(self.wb[real])
            for f in self._cache[real][:filas_max]:
                cab = [norm(x) for x in f]
                if all(any(r == c or c.startswith(r) for c in cab) for r in req):
                    return self._cache[real]
        return None

    def tabla(self, nombre: str, requeridas: tuple[str, ...]) -> list[dict[str, Any]]:
        """Hoja tipo tabla: localiza la fila de encabezados que contiene `requeridas`."""
        filas = self.hoja(nombre) or []
        req = [norm(r) for r in requeridas]
        for i, fila in enumerate(filas[:30]):
            cab = [norm(x) for x in fila]
            if all(r in cab for r in req):
                out = []
                for f in filas[i + 1:]:
                    if not any(x not in (None, "") for x in f):
                        continue
                    out.append({cab[j]: f[j] for j in range(min(len(cab), len(f))) if cab[j]})
                return out
        return []

    def datos(self) -> dict[str, Any]:
        out = {}
        for f in self.hoja("Datos") or []:
            if f and f[0] is not None and len(f) > 1:
                out[norm(f[0])] = f[1]
        return out

    def prueba(self) -> dict[str, dict[str, float]]:
        """Balance de prueba: {cuenta: {saldo_anterior, debitos, creditos, nuevo_saldo, nombre}}."""
        filas = (self.hoja("Prueba") or self.hoja_con(("debitos", "creditos")) or self.hoja_con(("debito", "credito"))
                 or self.hoja_con(("debe", "haber")) or [])
        out: dict[str, dict[str, float]] = {}
        icol = None
        for f in filas:
            cab = [norm(x) for x in f]
            if icol is None:
                def col(*opciones: str) -> int | None:
                    return next((cab.index(o) for o in opciones if o in cab), None)
                deb, cre = col("debitos", "debito", "debe", "movimiento debito"), col("creditos", "credito", "haber",
                                                                                      "movimiento credito")
                if deb is not None and cre is not None:
                    icol = {"debitos": deb, "creditos": cre}
                    ant = col("saldo anterior", "saldo inicial", "saldo ant")
                    nue = col("nuevo saldo", "saldo final", "saldo actual", "saldo nuevo")
                    icol["anterior"] = ant if ant is not None else deb - 1
                    icol["nuevo"] = nue if nue is not None else cre + 1
                continue
            # columnas típicas: nivel (número), cuenta (texto), descripción…
            cuenta = next((x.strip() for x in f[:3] if isinstance(x, str) and re.fullmatch(r"\d{1,12}", x.strip())), None)
            if not cuenta:
                nums = [x for x in f[:2] if _es_num(x) and float(x).is_integer()]
                if len(nums) == 2:
                    cuenta = str(int(nums[1]))
                else:
                    continue
            def num(i: int) -> float:
                v = f[i] if i < len(f) else 0
                return float(v) if _es_num(v) else 0.0
            nombre = next((str(x) for x in f[:4] if isinstance(x, str) and not re.fullmatch(r"\d+", x.strip())), "")
            if cuenta in out:  # algunas hojas repiten un bloque más abajo (otro corte): manda el primero
                continue
            out[cuenta] = {"saldo_anterior": num(icol["anterior"]), "debitos": num(icol["debitos"]),
                           "creditos": num(icol["creditos"]), "nuevo_saldo": num(icol["nuevo"]), "nombre": nombre}
        return out

    def auxiliar(self, prefijos: tuple[str, ...]) -> list[dict[str, Any]]:
        """Movimientos del auxiliar de las cuentas que empiezan por `prefijos`."""
        filas = self.hoja("Auxiliar") or self.hoja_con(("cuenta", "fecha")) or []
        cab = None
        out = []
        for f in filas:
            if cab is None:
                c = [norm(x) for x in f]
                if "cuenta" in c and "fecha" in c:
                    cab = c
                continue
            fila = {cab[j]: f[j] for j in range(min(len(cab), len(f))) if cab[j]}
            cuenta = str(fila.get("cuenta") or "").strip()
            if not cuenta.startswith(prefijos):
                continue
            fecha = fila.get("fecha")
            if isinstance(fecha, datetime):
                fecha = fecha.date()
            elif isinstance(fecha, str):
                try:
                    fecha = date.fromisoformat(fecha[:10])
                except ValueError:
                    fecha = None
            saldo = fila.get("saldo")
            out.append({"cuenta": cuenta, "fecha": fecha, "nit": fila.get("nit"), "nombre": fila.get("nombre"),
                        "comprobante": fila.get("comprobante"), "detalle": str(fila.get("detalle") or ""),
                        "valor": float(saldo) if _es_num(saldo) else 0.0})
        return out

    def cerrar(self) -> None:
        try:
            self.wb.close()
        except Exception:
            pass
